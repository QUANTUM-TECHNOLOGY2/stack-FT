"""
================================================================================
GESTIONNAIRE DE FICHES TECHNIQUES - SUPABASE EDITION
================================================================================
Fichier unique avec :
1. Base de données Supabase (PostgreSQL)
2. Authentification JWT
3. CRUD complet des fiches techniques
4. Visualisation en ligne (PDF, images, documents)
5. Versioning des fichiers
6. Système de tags
7. Notifications
8. Statistiques
9. Design professionnel avec couleurs #D61C4E et #293462
10. Drag & drop pour l'import
================================================================================
"""

import os
import sys
import uuid
import json
import io
import base64
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlparse
import requests
from dotenv import load_dotenv
import sqlite3


# Simple response wrapper used by the remote Supabase helpers
class RemoteSupabaseResponse:
    def __init__(self, data=None, count=None, status_code=0, error=None):
        self.data = data if data is not None else []
        self.count = count
        self.status_code = status_code
        self.error = error
    def __repr__(self):
        return f"RemoteSupabaseResponse(data={self.data!r}, count={self.count!r}, status_code={self.status_code!r}, error={self.error!r})"

try:
    import magic
except Exception:  # pragma: no cover
    magic = None

# Third-party imports
import PyPDF2
try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None
import bcrypt

try:
    from supabase import create_client, Client
except Exception:  # pragma: no cover
    create_client = None
    Client = object
    class RemoteSupabaseTable:
        def __init__(self, client, table_name):
            self.client = client
            self.table_name = table_name
            self._filters = []
            self._action = None
            self._select = '*'
            self._order = None
            self._limit = None
            self._count = False
            self._range = None
            self._data = None

        def in_(self, column, values):
            if isinstance(values, (list, tuple)):
                formatted_values = []
                for v in values:
                    if isinstance(v, str):
                        formatted_values.append(f'"{v}"')
                    else:
                        formatted_values.append(str(v))
                values = ','.join(formatted_values)
            self._filters.append((column, 'in', f'({values})'))
            return self

        def select(self, select_str='*'):
            self._action = 'select'
            self._select = select_str
            return self

        def eq(self, column, value):
            self._filters.append((column, 'eq', value))
            return self

        def order(self, column, desc=False):
            self._order = f"{column}.desc" if desc else f"{column}.asc"
            return self

        def limit(self, n):
            try:
                self._limit = int(n)
            except Exception:
                self._limit = n
            return self

        def insert(self, data):
            self._action = 'insert'
            self._data = data
            return self

        def update(self, data):
            self._action = 'update'
            self._data = data
            return self

        def delete(self):
            self._action = 'delete'
            return self

        def or_(self, expr):
            # PostgREST uses 'or' query param with comma-separated expressions
            self._filters.append(('or', 'or', expr))
            return self

        def range(self, start, end):
            self._range = (start, end)
            return self

        def filter(self, column, op, value):
            self._filters.append((column, op, value))
            return self

        def _build_params(self):
            params = {}
            if self._action == 'select':
                params['select'] = self._select
            if self._filters:
                for column, op, value in self._filters:
                    if column == 'or':
                        params['or'] = value
                    else:
                        if isinstance(value, bool):
                            value = str(value).lower()
                        params[column] = f"{op}.{value}"
            if self._order:
                params['order'] = self._order
            if self._limit is not None:
                params['limit'] = str(self._limit)
            return params

        def execute(self):
            url = f"{self.client.url}/rest/v1/{self.table_name}"
            headers = {
                'apikey': self.client.service_key,
                'Authorization': f'Bearer {self.client.service_key}',
            }

            try:
                if self._action == 'select':
                    if hasattr(self, '_range') and self._range is not None:
                        headers['Range'] = f"items={self._range[0]}-{self._range[1]}"
                    elif self._count:
                        headers['Range'] = 'items=0-0'
                    if self._count:
                        headers['Prefer'] = 'count=exact'
                    response = requests.get(url, headers=headers, params=self._build_params(), timeout=30)
                elif self._action == 'insert':
                    headers['Content-Type'] = 'application/json'
                    # Use Prefer header for PostgREST to request representation return
                    headers['Prefer'] = 'return=representation'
                    response = requests.post(url, headers=headers, json=self._data, timeout=30)
                    if response.status_code == 400 and response.text:
                        lowered = response.text.lower()
                        if 'return=representation' in lowered or 'failed to parse filter' in lowered or 'pgrst100' in lowered:
                            headers.pop('Prefer', None)
                            response = requests.post(url, headers=headers, json=self._data, timeout=30)
                elif self._action == 'update':
                    headers['Content-Type'] = 'application/json'
                    response = requests.patch(url, headers=headers, json=self._data, params=self._build_params(), timeout=30)
                elif self._action == 'delete':
                    response = requests.delete(url, headers=headers, params=self._build_params(), timeout=30)
                else:
                    response = requests.get(url, headers=headers, params=self._build_params(), timeout=30)

                if response.status_code >= 400:
                    err_text = response.text
                    try:
                        parsed = response.json()
                        err_text = parsed.get('message') or parsed.get('error') or err_text
                    except Exception:
                        pass
                    return RemoteSupabaseResponse(data=[], count=0, status_code=response.status_code, error=err_text)

                data = response.json() if response.text else []
                count = None
                if 'content-range' in response.headers:
                    content_range = response.headers['content-range']
                    if '/' in content_range:
                        total = content_range.split('/')[-1]
                        if total.isdigit():
                            count = int(total)
                return RemoteSupabaseResponse(data=data, count=count, status_code=response.status_code)
            except Exception:
                return RemoteSupabaseResponse(data=[], count=0, status_code=0, error='exception')


class RemoteSupabaseAuth:
    def __init__(self, client):
        self.client = client

    def _parse_error(self, response, data):
        if not isinstance(data, dict):
            return response.text or 'Erreur Supabase inconnue'
        return (
            data.get('msg')
            or data.get('error_description')
            or data.get('error')
            or data.get('message')
            or data.get('details')
            or response.text
        )

    def sign_in_with_password(self, credentials):
        url = f"{self.client.url}/auth/v1/token?grant_type=password"
        headers = {
            'apikey': self.client.key,
            'Content-Type': 'application/json',
        }
        response = requests.post(url, headers=headers, json=credentials, timeout=30)
        data = response.json() if response.text else {}
        if response.status_code >= 400:
            error = self._parse_error(response, data)
            return SimpleNamespace(user=None, session=None, error=error, status_code=response.status_code)
        session = SimpleNamespace(
            access_token=data.get('access_token'),
            refresh_token=data.get('refresh_token'),
        )
        user_data = data.get('user') or {}
        user = SimpleNamespace(**user_data) if user_data else None
        return SimpleNamespace(user=user, session=session, error=None, status_code=response.status_code)

    def sign_up(self, credentials):
        url = f"{self.client.url}/auth/v1/signup"
        headers = {
            'apikey': self.client.key,
            'Content-Type': 'application/json',
        }
        response = requests.post(url, headers=headers, json=credentials, timeout=30)
        data = response.json() if response.text else {}
        if response.status_code >= 400:
            if data.get('error_code') == 'over_email_send_rate_limit':
                return self._admin_sign_up(credentials)
            error = self._parse_error(response, data)
            # If Supabase rejects the email as invalid, attempt to create the user
            # via the admin endpoint (service key) which accepts broader values.
            try:
                err_text = (error or '').lower() if isinstance(error, str) else ''
                if 'invalid' in err_text or 'email address' in err_text or 'invalid email' in err_text:
                    return self._admin_sign_up(credentials)
            except Exception:
                pass
            return SimpleNamespace(user=None, session=None, error=error, status_code=response.status_code)
        session = SimpleNamespace(
            access_token=data.get('access_token'),
            refresh_token=data.get('refresh_token'),
        )
        user_data = None
        if isinstance(data, dict):
            if data.get('user'):
                user_data = data.get('user')
            elif isinstance(data.get('session'), dict) and data.get('session', {}).get('user'):
                user_data = data.get('session', {}).get('user')
            elif data.get('id') and data.get('email'):
                # Supabase may return the user object at the top level for signup
                user_data = data
        if user_data is None:
            user_data = {}
        user = SimpleNamespace(**user_data) if user_data else None
        return SimpleNamespace(user=user, session=session, error=None, status_code=response.status_code)

    def _admin_sign_up(self, credentials):
        admin_url = f"{self.client.url}/auth/v1/admin/users"
        headers = {
            'apikey': self.client.service_key,
            'Authorization': f'Bearer {self.client.service_key}',
            'Content-Type': 'application/json',
        }
        payload = {
            'email': credentials.get('email'),
            'password': credentials.get('password'),
            'email_confirm': True,
        }
        response = requests.post(admin_url, headers=headers, json=payload, timeout=30)
        if response.status_code >= 400:
            return SimpleNamespace(user=None, session=None)
        user_data = response.json() if response.text else {}
        user = SimpleNamespace(**user_data) if user_data else None
        if not user:
            return SimpleNamespace(user=None, session=None)
        login_response = self.sign_in_with_password({'email': credentials.get('email'), 'password': credentials.get('password')})
        if login_response.user:
            return login_response
        return SimpleNamespace(user=user, session=SimpleNamespace(access_token=None, refresh_token=None))

    def get_user(self, token):
        if not token:
            return None
        url = f"{self.client.url}/auth/v1/user"
        headers = {
            'apikey': self.client.key,
            'Authorization': f'Bearer {token}',
        }
        response = requests.get(url, headers=headers, timeout=30)
        if response.status_code >= 400:
            return None
        data = response.json() if response.text else {}
        # Supabase may return the user object directly or nested; normalize
        user_data = data.get('user') if isinstance(data, dict) and 'user' in data else data
        user = SimpleNamespace(**user_data) if user_data else None
        return SimpleNamespace(user=user)


def normalize_supabase_auth_response(response):
    if response is None:
        return {'user': None, 'session': None, 'error': None}
    user = None
    session = None
    error = None

    if isinstance(response, dict):
        user = response.get('user') or (response.get('session') or {}).get('user')
        session = response.get('session')
        error = response.get('error') or response.get('message') or response.get('msg')
    else:
        user = getattr(response, 'user', None)
        session = getattr(response, 'session', None)
        error = getattr(response, 'error', None)
        if user is None and session is not None:
            if isinstance(session, dict):
                user = session.get('user')
            else:
                user = getattr(session, 'user', None)
                if user is None and hasattr(session, 'id') and hasattr(session, 'email'):
                    user = session

    if isinstance(user, dict):
        user = SimpleNamespace(**user)
    return {'user': user, 'session': session, 'error': error}


def admin_signup_fallback(email, password):
    if not SUPABASE_URL or not SUPABASE_SERVICE_KEY:
        return None
    url = f"{SUPABASE_URL.rstrip('/')}/auth/v1/admin/users"
    headers = {
        'apikey': SUPABASE_SERVICE_KEY,
        'Authorization': f'Bearer {SUPABASE_SERVICE_KEY}',
        'Content-Type': 'application/json',
    }
    payload = {
        'email': email,
        'password': password,
        'email_confirm': True,
    }
    try:
        response = requests.post(url, headers=headers, json=payload, timeout=30)
    except Exception:
        return None
    if response.status_code >= 400:
        return None
    data = response.json() if response.text else {}
    return SimpleNamespace(**data) if isinstance(data, dict) and data else None


class RemoteSupabaseBucket:
    def __init__(self, client, bucket):
        self.client = client
        self.bucket = bucket

    def upload(self, path, file_data, options=None):
        endpoint = f"{self.client.url}/storage/v1/object/{self.bucket}/{path}"
        headers = {
            'apikey': self.client.service_key,
            'Authorization': f'Bearer {self.client.service_key}',
        }
        if options and 'content-type' in options:
            headers['Content-Type'] = options['content-type']
        response = requests.put(endpoint, data=file_data, headers=headers, timeout=60)
        return {'status_code': response.status_code, 'data': response.text}

    def get_public_url(self, path):
        return f"{self.client.url}/storage/v1/object/public/{self.bucket}/{path}"

    def remove(self, paths):
        if not paths:
            return {'success': False}
        endpoint = f"{self.client.url}/storage/v1/object/{self.bucket}/{paths[0]}"
        headers = {
            'apikey': self.client.service_key,
            'Authorization': f'Bearer {self.client.service_key}',
        }
        response = requests.delete(endpoint, headers=headers, timeout=30)
        return {'success': response.status_code in (200, 204)}


class RemoteSupabaseStorage:
    def __init__(self, client):
        self.client = client

    def from_(self, bucket):
        return RemoteSupabaseBucket(self.client, bucket)


class RemoteSupabaseClient:
    def __init__(self, url, key, service_key):
        self.url = url.rstrip('/')
        self.key = key
        self.service_key = service_key
        self.auth = RemoteSupabaseAuth(self)
        self.storage = RemoteSupabaseStorage(self)

    def table(self, name):
        return RemoteSupabaseTable(self, name)


# Django imports
import django
from django.core.management import execute_from_command_line
from django.core.wsgi import get_wsgi_application
from django.conf import settings
from django.conf.urls import include
from django.core.paginator import Paginator
from django.db import models
from django.db.models import Q, Count
from django.http import JsonResponse, HttpResponse, FileResponse, Http404
from django.shortcuts import render, get_object_or_404, redirect
from django.template import loader
from django.template.response import TemplateResponse
from django.urls import path, re_path
from django.utils import timezone
from django.utils.html import escape
from django.views.decorators.csrf import csrf_exempt, ensure_csrf_cookie
from django.views.decorators.http import require_http_methods
from django.views.generic import View, TemplateView

# =================================================================================
# CONFIGURATION DJANGO
# =================================================================================

BASE_DIR = Path(__file__).resolve().parent
STATIC_ROOT = BASE_DIR / 'staticfiles'
MEDIA_ROOT = BASE_DIR / 'media'
MEDIA_URL = '/media/'

load_dotenv(BASE_DIR / '.env')

# Configuration Supabase - À remplacer par vos identifiants
SUPABASE_URL = os.environ.get('SUPABASE_URL', 'https://votre-projet.supabase.co')
SUPABASE_KEY = os.environ.get('SUPABASE_PUBLISHABLE_KEY') or os.environ.get('SUPABASE_KEY') or os.environ.get('SUPABASE_SECRET_KEY')
SUPABASE_SERVICE_KEY = os.environ.get('SUPABASE_SECRET_KEY') or os.environ.get('SUPABASE_SERVICE_KEY')

# Allow switching to a local SQLite backend for dev/testing using USE_LOCAL_BACKEND=1
USE_LOCAL_BACKEND = os.environ.get('USE_LOCAL_BACKEND') in ('1', 'true', 'True')
if USE_LOCAL_BACKEND:
    from local_backend import create_local_backend
    print('Using Local backend (SQLite) for auth and storage')
    supabase = create_local_backend(BASE_DIR, MEDIA_ROOT)
else:
    # Initialisation du client Supabase distant uniquement
    print('Using Supabase remote-only backend')
    if create_client is not None:
        try:
            supabase = create_client(SUPABASE_URL, SUPABASE_KEY)
            print('Supabase python client initialized successfully')
        except Exception as exc:
            print('Supabase python client init failed, using HTTP fallback:', exc)
            supabase = RemoteSupabaseClient(SUPABASE_URL, SUPABASE_KEY, SUPABASE_SERVICE_KEY)
    else:
        print('Supabase python client unavailable, using HTTP fallback')
        supabase = RemoteSupabaseClient(SUPABASE_URL, SUPABASE_KEY, SUPABASE_SERVICE_KEY)

SETTINGS_DICT = {
    'DEBUG': True,
    'SECRET_KEY': 'django-insecure-supabase-key-2024',
    'ALLOWED_HOSTS': ['*'],
    'X_FRAME_OPTIONS': 'SAMEORIGIN',
    'ROOT_URLCONF': __name__,
    'WSGI_APPLICATION': f'{__name__}.application',
    'INSTALLED_APPS': [
        'django.contrib.admin',
        'django.contrib.auth',
        'django.contrib.contenttypes',
        'django.contrib.sessions',
        'django.contrib.messages',
        'django.contrib.staticfiles',
        'rest_framework',
        'rest_framework_simplejwt',
    ],
    'MIDDLEWARE': [
        'django.middleware.security.SecurityMiddleware',
        'django.contrib.sessions.middleware.SessionMiddleware',
        'django.middleware.common.CommonMiddleware',
        'django.middleware.csrf.CsrfViewMiddleware',
        'django.contrib.auth.middleware.AuthenticationMiddleware',
        'django.contrib.messages.middleware.MessageMiddleware',
        'django.middleware.clickjacking.XFrameOptionsMiddleware',
    ],
    'TEMPLATES': [{
        'BACKEND': 'django.template.backends.django.DjangoTemplates',
        'DIRS': [BASE_DIR / 'templates'],
        'APP_DIRS': True,
        'OPTIONS': {
            'context_processors': [
                'django.template.context_processors.debug',
                'django.template.context_processors.request',
                'django.contrib.auth.context_processors.auth',
                'django.contrib.messages.context_processors.messages',
            ],
        },
    }],
    'DATABASES': {
        'default': {
            'ENGINE': 'django.db.backends.sqlite3',
            'NAME': BASE_DIR / 'db_local.sqlite3',
        }
    },
    'SESSION_ENGINE': 'django.contrib.sessions.backends.signed_cookies',
    'SESSION_COOKIE_AGE': 1209600,
    'SESSION_SAVE_EVERY_REQUEST': True,
    'STATIC_URL': '/static/',
    'STATIC_ROOT': STATIC_ROOT,
    'STATICFILES_DIRS': [BASE_DIR / 'static'],
    'MEDIA_ROOT': MEDIA_ROOT,
    'MEDIA_URL': MEDIA_URL,
    'LANGUAGE_CODE': 'fr-fr',
    'TIME_ZONE': 'Africa/Porto-Novo',
    'USE_I18N': True,
    'USE_TZ': True,
    'DEFAULT_AUTO_FIELD': 'django.db.models.BigAutoField',
    'LOGIN_URL': '/login/',
    'LOGIN_REDIRECT_URL': '/dashboard/',
    'LOGOUT_REDIRECT_URL': '/login/',
    'REST_FRAMEWORK': {
        'DEFAULT_AUTHENTICATION_CLASSES': [
            'rest_framework_simplejwt.authentication.JWTAuthentication',
        ],
        'DEFAULT_PERMISSION_CLASSES': [
            'rest_framework.permissions.IsAuthenticated',
        ],
        'DEFAULT_PAGINATION_CLASS': 'rest_framework.pagination.PageNumberPagination',
        'PAGE_SIZE': 20,
    },
    'SIMPLE_JWT': {
        'ACCESS_TOKEN_LIFETIME': timedelta(days=1),
        'REFRESH_TOKEN_LIFETIME': timedelta(days=7),
    },
}

for _key, _value in SETTINGS_DICT.items():
    globals()[_key] = _value

settings.configure(**SETTINGS_DICT)
django.setup()

from rest_framework import viewsets, status, filters
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.response import Response
from rest_framework.serializers import ModelSerializer
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework_simplejwt.authentication import JWTAuthentication
from rest_framework_simplejwt.views import TokenObtainPairView, TokenRefreshView
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser

from django.contrib import admin
from django.contrib.auth.models import AbstractUser, Group
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.hashers import make_password, check_password
from django.contrib.auth.middleware import AuthenticationMiddleware
from django.contrib.messages.middleware import MessageMiddleware
from django.contrib.sessions.middleware import SessionMiddleware
from django.contrib.staticfiles.urls import staticfiles_urlpatterns
from django.conf.urls.static import static

# Création des dossiers
os.makedirs(BASE_DIR / 'templates', exist_ok=True)
os.makedirs(BASE_DIR / 'static', exist_ok=True)
os.makedirs(MEDIA_ROOT / 'uploads', exist_ok=True)

# =================================================================================
# MODÈLES SUPABASE
# =================================================================================

# Tables Supabase à créer (SQL)
SUPABASE_SCHEMA = """
-- Table des utilisateurs (gérée par Supabase Auth)
-- Table des profils utilisateurs
CREATE TABLE IF NOT EXISTS public.profiles (
    id UUID REFERENCES auth.users(id) PRIMARY KEY,
    username TEXT UNIQUE NOT NULL,
    full_name TEXT,
    avatar_url TEXT,
    role TEXT DEFAULT 'user',
    created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW(),
    updated_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
);

-- Table des tags
CREATE TABLE IF NOT EXISTS public.tags (
    id UUID DEFAULT gen_random_uuid() PRIMARY KEY,
    name TEXT UNIQUE NOT NULL,
    color TEXT DEFAULT '#293462',
    description TEXT,
    created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
);

-- Table des fiches techniques
CREATE TABLE IF NOT EXISTS public.fiches (
    id UUID DEFAULT gen_random_uuid() PRIMARY KEY,
    reference TEXT UNIQUE NOT NULL,
    name TEXT NOT NULL,
    description TEXT,
    category TEXT,
    manufacturer TEXT,
    version TEXT DEFAULT '1.0',
    file_url TEXT,
    file_name TEXT,
    file_size INTEGER,
    file_type TEXT,
    file_preview JSONB,
    author_id UUID REFERENCES auth.users(id),
    created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW(),
    updated_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
);

-- Table de liaison fiches-tags
CREATE TABLE IF NOT EXISTS public.fiche_tags (
    fiche_id UUID REFERENCES public.fiches(id) ON DELETE CASCADE,
    tag_id UUID REFERENCES public.tags(id) ON DELETE CASCADE,
    PRIMARY KEY (fiche_id, tag_id)
);

-- Table des versions
CREATE TABLE IF NOT EXISTS public.versions (
    id UUID DEFAULT gen_random_uuid() PRIMARY KEY,
    fiche_id UUID REFERENCES public.fiches(id) ON DELETE CASCADE,
    version TEXT NOT NULL,
    file_url TEXT,
    file_name TEXT,
    comment TEXT,
    author_id UUID REFERENCES auth.users(id),
    created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
);

-- Table des notifications
CREATE TABLE IF NOT EXISTS public.notifications (
    id UUID DEFAULT gen_random_uuid() PRIMARY KEY,
    recipient_id UUID REFERENCES auth.users(id) ON DELETE CASCADE,
    type TEXT NOT NULL,
    message TEXT NOT NULL,
    fiche_id UUID REFERENCES public.fiches(id) ON DELETE CASCADE,
    is_read BOOLEAN DEFAULT FALSE,
    created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
);

-- Index pour les performances
CREATE INDEX IF NOT EXISTS idx_fiches_author ON public.fiches(author_id);
CREATE INDEX IF NOT EXISTS idx_fiches_category ON public.fiches(category);
CREATE INDEX IF NOT EXISTS idx_fiches_created ON public.fiches(created_at DESC);
CREATE INDEX IF NOT EXISTS idx_notifications_recipient ON public.notifications(recipient_id, is_read);
CREATE INDEX IF NOT EXISTS idx_notifications_created ON public.notifications(created_at DESC);
"""

# =================================================================================
# UTILITAIRES SUPABASE
# =================================================================================

def supabase_upload_file(file_data, filename):
    """Upload un fichier vers Supabase Storage"""
    try:
        # Générer un nom unique
        unique_name = f"{uuid.uuid4().hex[:8]}_{filename}"
        file_path = f"uploads/{unique_name}"
        content_type = guess_mime_from_extension(filename) or 'application/octet-stream'

        # Upload vers Supabase Storage
        response = supabase.storage.from_('fiches').upload(
            file_path,
            file_data,
            {'content-type': content_type}
        )

        public_url = None
        error = None
        if response:
            if isinstance(response, dict):
                error = response.get('error') or response.get('message')
            elif hasattr(response, 'error'):
                error = getattr(response, 'error', None)
            public_result = supabase.storage.from_('fiches').get_public_url(file_path)
            if isinstance(public_result, dict):
                public_url = public_result.get('public_url') or public_result.get('publicUrl')
                if not public_url:
                    data = public_result.get('data')
                    if isinstance(data, dict):
                        public_url = data.get('public_url') or data.get('publicUrl')
            elif hasattr(public_result, 'data') and isinstance(public_result.data, dict):
                public_url = public_result.data.get('public_url') or public_result.data.get('publicUrl')
            elif isinstance(public_result, str):
                public_url = public_result
            if not public_url:
                public_url = str(public_result)

        try:
            local_dir = BASE_DIR / 'media' / 'uploads'
            local_dir.mkdir(parents=True, exist_ok=True)
            local_path = local_dir / unique_name
            with open(local_path, 'wb') as lf:
                lf.write(file_data)
        except Exception:
            local_path = None

        if public_url and not error:
            return {
                'success': True,
                'url': public_url,
                'path': file_path,
                'name': filename,
                'local_path': str(local_path) if local_path else None,
                'local_url': f'/media/uploads/{unique_name}' if local_path else None,
            }

        if local_path:
            return {
                'success': True,
                'url': f'/media/uploads/{unique_name}',
                'path': file_path,
                'name': filename,
                'local_path': str(local_path),
                'local_url': f'/media/uploads/{unique_name}',
                'warning': 'Supabase storage unavailable, using local media fallback',
                'upload_error': error,
            }

        return {'success': False, 'error': error or 'Upload failed'}
    except Exception as e:
        return {'success': False, 'error': str(e)}

def extract_storage_path(file_url):
    if not file_url:
        return None
    try:
        parsed = urlparse(file_url)
        path = parsed.path or file_url
        markers = [
            '/storage/v1/object/public/fiches/uploads/',
            '/storage/v1/object/public/fiches/',
            '/storage/v1/object/fiches/uploads/',
            '/storage/v1/object/fiches/',
        ]
        for marker in markers:
            if marker in path:
                return path.split(marker, 1)[1]
        if path.startswith('/uploads/'):
            return path[len('/uploads/'):]
        if path.startswith('/fiches/'):
            return path[len('/fiches/'):]
        return path.lstrip('/')
    except Exception:
        return None


def build_supabase_storage_download_url(file_url):
    if not file_url:
        return None
    try:
        parsed = urlparse(file_url)
        base_host = urlparse(SUPABASE_URL).netloc if SUPABASE_URL else None
        if base_host and parsed.netloc.endswith(base_host):
            path = parsed.path or ''
            if '/storage/v1/object/public/' in path:
                return f"{SUPABASE_URL.rstrip('/')}/storage/v1/object/{path.split('/storage/v1/object/public/', 1)[1]}"
            if '/storage/v1/object/' in path:
                return f"{SUPABASE_URL.rstrip('/')}/storage/v1/object/{path.split('/storage/v1/object/', 1)[1]}"
        return file_url
    except Exception:
        return file_url


def build_supabase_storage_download_urls(file_url):
    urls = []
    if not file_url:
        return urls
    urls.append(file_url)
    alt_url = build_supabase_storage_download_url(file_url)
    if alt_url and alt_url != file_url:
        urls.append(alt_url)

    parsed = urlparse(file_url)
    if parsed.scheme and parsed.netloc:
        path = parsed.path or ''
        if '/storage/v1/object/public/fiches/' in path and '/uploads/' not in path:
            urls.append(f"{SUPABASE_URL.rstrip('/')}/storage/v1/object/public/fiches/uploads/{path.split('/storage/v1/object/public/fiches/',1)[1]}")
        if '/storage/v1/object/fiches/' in path and '/uploads/' not in path:
            urls.append(f"{SUPABASE_URL.rstrip('/')}/storage/v1/object/fiches/uploads/{path.split('/storage/v1/object/fiches/',1)[1]}")
    return list(dict.fromkeys(urls))


def download_remote_file(file_url):
    if not file_url:
        return None, 'Missing file_url', None
    headers = {
        'apikey': SUPABASE_SERVICE_KEY,
        'Authorization': f'Bearer {SUPABASE_SERVICE_KEY}',
    }
    tried = []
    for url in build_supabase_storage_download_urls(file_url):
        try:
            response = requests.get(url, headers=headers, timeout=30)
            tried.append((url, response.status_code, response.text[:200]))
            if response.status_code == 200:
                return response.content, None, response.headers.get('content-type')
        except Exception as e:
            tried.append((url, 'exception', str(e)))
    details = '; '.join([f"{u} -> {s}" for u, s, _ in tried])
    return None, f'No remote file could be fetched: {details}', None


def find_local_fallback_path(file_url=None, file_name=None):
    if not file_url and not file_name:
        return None
    filename = None
    if file_name:
        filename = Path(file_name).name
    else:
        try:
            parsed = urlparse(file_url)
            filename = Path(parsed.path).name
        except Exception:
            filename = None
    if not filename:
        return None
    local_dir = BASE_DIR / 'media' / 'uploads'
    exact_path = local_dir / filename
    if exact_path.exists():
        return exact_path

    # Cherche un fichier local comportant le même suffixe ou le même nom sans préfixe
    target_suffix = filename
    if '_' in filename:
        target_suffix = filename.split('_', 1)[1]
    for local_file in local_dir.iterdir():
        if not local_file.is_file():
            continue
        if local_file.name == filename:
            return local_file
        if local_file.name.endswith(target_suffix):
            return local_file
        if filename in local_file.name:
            return local_file
    return None


def supabase_delete_file(file_path):
    """Supprime un fichier de Supabase Storage"""
    if not file_path:
        return False
    try:
        supabase.storage.from_('fiches').remove([file_path])
        return True
    except Exception:
        return False


def get_user_profile(user_id):
    """Récupère le profil d'un utilisateur"""
    try:
        response = supabase.table('profiles').select('*').eq('id', user_id).execute()
        if response.data:
            return response.data[0]
        return None
    except Exception:
        return None


def get_profile_display_name(profile=None, user_id=None):
    """Retourne un nom d'affichage lisible pour un profil."""
    if profile is None and user_id:
        profile = get_user_profile(user_id)
    if isinstance(profile, dict):
        full_name = (profile.get('full_name') or '').strip()
        username = (profile.get('username') or '').strip()
        return full_name or username or 'Inconnu'
    return 'Inconnu'


def upsert_user_profile(user_id, username='', full_name='', role='user'):
    """Crée ou met à jour le profil utilisateur."""
    try:
        existing = get_user_profile(user_id)
        profile_data = {
            'id': user_id,
            'username': username or '',
            'full_name': full_name or '',
            'role': role,
        }
        if existing:
            return supabase.table('profiles').update(profile_data).eq('id', user_id).execute()
        return supabase.table('profiles').insert(profile_data).execute()
    except Exception:
        return None


def create_notification_supabase(recipient_id, type, message, fiche_id=None):
    """Crée une notification dans Supabase"""
    try:
        supabase.table('notifications').insert({
            'recipient_id': recipient_id,
            'type': type,
            'message': message,
            'fiche_id': fiche_id
        }).execute()
    except Exception as e:
        print(f"Erreur notification: {e}")

# =================================================================================
# DECORATEURS
# =================================================================================

def login_required_api(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        token = request.headers.get('Authorization', '').replace('Bearer ', '')
        if token:
            try:
                user = supabase.auth.get_user(token)
                if user:
                    request.user_id = user.user.id
                    request.user_email = user.user.email
                    return view_func(request, *args, **kwargs)
            except Exception:
                pass

        user_id = request.session.get('user_id')
        if user_id:
            request.user_id = user_id
            request.user_email = request.session.get('user_email', '')
            return view_func(request, *args, **kwargs)

        return JsonResponse({'error': 'Authentification requise'}, status=401)
    return wrapper


def require_app_login(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        user_id = request.session.get('user_id')
        if not user_id:
            if request.headers.get('x-requested-with') == 'XMLHttpRequest' or request.path.startswith('/api/'):
                return JsonResponse({'error': 'Authentification requise'}, status=401)
            return redirect('login')

        request.user = SimpleNamespace(
            id=user_id,
            username=request.session.get('username', ''),
            email=request.session.get('user_email', ''),
            full_name=request.session.get('full_name', ''),
            is_authenticated=True,
            is_staff=False,
            get_full_name=lambda: request.session.get('full_name', ''),
        )
        request.user_id = user_id
        request.user_email = request.session.get('user_email', '')
        return view_func(request, *args, **kwargs)
    return wrapper

# =================================================================================
# VUES FRONTEND
# =================================================================================

@ensure_csrf_cookie
def login_view(request):
    if request.method == 'POST':
        email = request.POST.get('email', '')
        password = request.POST.get('password', '')
        
        # Log d'entrée
        log_file = open('login_debug.log', 'a')
        log_file.write(f"\n[LOGIN_VIEW] POST reçu - Email: {email}\n")
        log_file.close()
        
        try:
            log_file = open('login_debug.log', 'a')
            log_file.write(f"[LOGIN_VIEW] Appel supabase.auth.sign_in_with_password\n")
            log_file.close()
            
            response = supabase.auth.sign_in_with_password({'email': email, 'password': password})
            
            log_file = open('login_debug.log', 'a')
            log_file.write(f"[LOGIN_VIEW] Response reçue: {response}\n")
            log_file.write(f"[LOGIN_VIEW] Response.user: {getattr(response, 'user', None)}\n")
            log_file.write(f"[LOGIN_VIEW] Response.error: {getattr(response, 'error', None)}\n")
            log_file.close()
            
            if response.user:
                profile = get_user_profile(response.user.id)
                request.session['user_id'] = response.user.id
                request.session['user_email'] = response.user.email
                request.session['username'] = profile.get('username') if profile else email.split('@')[0]
                request.session['full_name'] = profile.get('full_name') if profile else ''
                request.session['role'] = profile.get('role') if profile else 'user'
                return redirect('dashboard')
            error = getattr(response, 'error', None) or 'Identifiants invalides'
            return render(request, 'login.html', {'error': error})
        except Exception as exc:
            log_file = open('login_debug.log', 'a')
            log_file.write(f"[LOGIN_VIEW] Exception: {exc}\n")
            log_file.close()
            pass
        return render(request, 'login.html', {'error': 'Identifiants invalides'})

    if request.session.get('user_id'):
        return redirect('dashboard')
    return render(request, 'login.html')

@ensure_csrf_cookie
def register_view(request):
    import sys
    log_file = open('c:\\Users\\kpeho\\Downloads\\Quantum\\register_debug.log', 'a')
    if request.method == 'POST':
        email = request.POST.get('email', '')
        password1 = request.POST.get('password1', '')
        password2 = request.POST.get('password2', '')
        username = request.POST.get('username', '').strip() or email.split('@')[0]
        full_name = request.POST.get('full_name', '').strip()
        
        log_file.write(f'\n[DEBUG] Inscription tentée pour: {email}\n')
        log_file.flush()
        
        if not full_name:
            log_file.write('[DEBUG] Nom complet manquant\n')
            log_file.close()
            return render(request, 'register.html', {'error': 'Le nom complet est obligatoire'})
        
        if password1 != password2:
            log_file.write(f'[DEBUG] Mots de passe ne correspondent pas\n')
            log_file.close()
            return render(request, 'register.html', {'error': 'Les mots de passe ne correspondent pas'})
        
        try:
            log_file.write(f'[DEBUG] Appel sign_up({email})\n')
            log_file.flush()
            try:
                response = supabase.auth.sign_up({'email': email, 'password': password1})
            except sqlite3.IntegrityError as ie:
                log_file.write(f'[DEBUG] IntegrityError during sign_up: {ie}\n')
                log_file.close()
                return render(request, 'register.html', {'error': "Cet e-mail est déjà utilisé."})
            log_file.write(f'[DEBUG] Response type: {type(response)}, has user: {hasattr(response, "user")}\n')
            if hasattr(response, 'user'):
                log_file.write(f'[DEBUG] response.user value: {getattr(response, "user", None)}\n')
            if hasattr(response, 'session'):
                log_file.write(f'[DEBUG] response.session value: {response.session}\n')
            if hasattr(response, 'error'):
                log_file.write(f'[DEBUG] response.error value: {response.error}\n')
            log_file.flush()

            normalized = normalize_supabase_auth_response(response)
            user = normalized['user']
            session = normalized['session']
            error_msg = normalized['error']

            if not user and session:
                access_token = getattr(session, 'access_token', None) if not isinstance(session, dict) else session.get('access_token')
                if access_token:
                    try:
                        sup_user = supabase.auth.get_user(access_token)
                        if sup_user and getattr(sup_user, 'user', None):
                            user = sup_user.user
                            log_file.write(f'[DEBUG] Obtained user from session token: {user}\n')
                    except Exception:
                        log_file.write('[DEBUG] get_user with access_token failed\n')

            if not user and (not error_msg or 'rate limit' in str(error_msg).lower() or 'email rate limit' in str(error_msg).lower()):
                log_file.write('[DEBUG] Attempting admin signup fallback due to missing user or rate limit\n')
                log_file.flush()
                fallback_user = admin_signup_fallback(email, password1)
                if fallback_user:
                    user = fallback_user
                    log_file.write(f'[DEBUG] Admin fallback created user: {getattr(user, "id", None)}\n')
                else:
                    log_file.write('[DEBUG] Admin fallback did not create a user\n')

            if not user and not error_msg:
                log_file.write('[DEBUG] No user and no error: trying sign_in_with_password fallback\n')
                log_file.flush()
                try:
                    signin = supabase.auth.sign_in_with_password({'email': email, 'password': password1})
                    log_file.write(f'[DEBUG] sign_in_with_password fallback: user={getattr(signin, "user", None)} error={getattr(signin, "error", None)} status={getattr(signin, "status_code", None)}\n')
                    if getattr(signin, 'user', None):
                        user = signin.user
                    elif getattr(signin, 'session', None) is not None:
                        fallback_access_token = getattr(signin.session, 'access_token', None) if not isinstance(signin.session, dict) else signin.session.get('access_token')
                        if fallback_access_token:
                            sup_user = supabase.auth.get_user(fallback_access_token)
                            if sup_user and getattr(sup_user, 'user', None):
                                user = sup_user.user
                                log_file.write(f'[DEBUG] Obtained user from signin fallback token: {user}\n')
                except Exception as exc:
                        log_file.write(f'[DEBUG] sign_in_with_password fallback exception: {exc}\n')

            if user:
                user_id = getattr(user, 'id', None)
                log_file.write(f'[DEBUG] Utilisateur créé: {user_id}\n')
                upsert_user_profile(user_id, username=username, full_name=full_name, role='user')
                log_file.close()
                request.session['user_id'] = user_id
                request.session['user_email'] = getattr(user, 'email', '')
                request.session['username'] = username
                request.session['full_name'] = full_name
                request.session['role'] = 'user'
                return redirect('dashboard')
            else:
                error_msg = getattr(response, 'error', None) or getattr(response, 'message', None) or getattr(response, 'details', None)
                error_msg = error_msg or 'Erreur lors de la création du compte'
                log_file.write(f'[DEBUG] ERREUR: response.user est None ou manquant, error={error_msg}\n')
                log_file.close()
                return render(request, 'register.html', {'error': error_msg})
        except Exception as exc:
            log_file.write(f'[DEBUG] Exception: {exc}\n')
            import traceback
            traceback.print_exc(file=log_file)
            log_file.close()
            return render(request, 'register.html', {'error': f'Erreur: {str(exc)}'})
    
    log_file.close()
    if request.session.get('user_id'):
        return redirect('dashboard')
    return render(request, 'register.html')


@ensure_csrf_cookie
def register_view_fixed(request):
    # GET: return a minimal, explicit HTML form to guarantee `required` attribute
    if request.method == 'GET':
        try:
            from django.middleware.csrf import get_token
            token = get_token(request)
        except Exception:
            token = ''
        html = f"""
<!doctype html>
<html><head><meta charset="utf-8"><title>Inscription</title></head><body>
<form method="post" action="/register/">
<input type="hidden" name="csrfmiddlewaretoken" value="{token}">
<label>Nom d'utilisateur *<input type="text" name="username" required></label><br>
<label>Email *<input type="email" name="email" required></label><br>
<label>Nom complet *<input type="text" name="full_name" required></label><br>
<label>Mot de passe *<input type="password" name="password1" required></label><br>
<label>Confirmer mot de passe *<input type="password" name="password2" required></label><br>
<button type="submit">S'inscrire</button>
</form>
</body></html>
"""
        return HttpResponse(html)
    # POST: delegate to existing register_view for full processing
    return register_view(request)

@require_app_login
def dashboard_view(request):
    context = {
        'page_title': 'Tableau de bord',
        'user': request.user,
    }
    return render(request, 'dashboard.html', context)

@require_app_login
def fiche_detail_view(request, fiche_id):
    try:
        response = supabase.table('fiches').select('*').eq('id', fiche_id).execute()
        if not response.data:
            return redirect('dashboard')
        fiche = response.data[0]
        profile = get_user_profile(fiche.get('author_id')) if fiche.get('author_id') else None
        fiche['profiles'] = {
            'username': profile.get('username') if profile else None,
            'full_name': profile.get('full_name') if profile else None,
        }
        fiche['author_display'] = get_profile_display_name(profile)
        # Récupérer les tags
        tags_response = supabase.table('fiche_tags').select('tags(*)').eq('fiche_id', fiche_id).execute()
        fiche['tags'] = [t['tags'] for t in tags_response.data] if tags_response.data else []
        # Récupérer les versions
        versions_response = supabase.table('versions').select('*').eq('fiche_id', fiche_id).order('created_at', desc=True).execute()
        versions = versions_response.data if versions_response.data else []
        if versions:
            version_author_ids = [v.get('author_id') for v in versions if v.get('author_id')]
            version_profiles = {}
            if version_author_ids:
                version_profiles_resp = supabase.table('profiles').select('id,username,full_name').in_('id', list(set(version_author_ids))).execute()
                version_profiles = {
                    profile['id']: profile for profile in (version_profiles_resp.data or []) if profile.get('id')
                }
            for version in versions:
                version_profile = version_profiles.get(version.get('author_id'))
                version['author_display'] = get_profile_display_name(version_profile)
        fiche['versions'] = versions
        
        context = {
            'fiche': fiche,
            'page_title': fiche.get('name', 'Fiche'),
        }
        return render(request, 'fiche_detail.html', context)
    except Exception:
        return redirect('dashboard')

def logout_view(request):
    request.session.flush()
    return redirect('login')


def debug_supabase_insert(request):
    """Endpoint de debug qui tente une insertion minimale dans la table `fiches`."""
    try:
        import requests as _requests
        url = f"{SUPABASE_URL.rstrip('/')}/rest/v1/fiches"
        headers = {
            'apikey': SUPABASE_SERVICE_KEY,
            'Authorization': f'Bearer {SUPABASE_SERVICE_KEY}',
            'Content-Type': 'application/json',
            'Prefer': 'return=representation'
        }
        payload = {'reference': f'TEST-{uuid.uuid4().hex[:8]}', 'name': 'Debug Insert'}
        r = _requests.post(url, headers=headers, json=payload, timeout=30)
        return HttpResponse(f"STATUS {r.status_code}\n{r.text}", content_type='text/plain')
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        return HttpResponse(f"EXCEPTION {e}\n{tb}", content_type='text/plain')

# =================================================================================
# VUES API
# =================================================================================

@csrf_exempt
@require_http_methods(["POST"])
def api_login(request):
    try:
        data = json.loads(request.body)
        email = data.get('email')
        password = data.get('password')
        
        # Authentification avec Supabase
        response = supabase.auth.sign_in_with_password({
            'email': email,
            'password': password
        })
        
        if response.user:
            # Récupérer le profil
            profile = get_user_profile(response.user.id)
            
            return JsonResponse({
                'access': response.session.access_token,
                'refresh': response.session.refresh_token,
                'user': {
                    'id': response.user.id,
                    'email': response.user.email,
                    'username': profile.get('username') if profile else email.split('@')[0],
                    'full_name': profile.get('full_name') if profile else '',
                    'role': profile.get('role') if profile else 'user',
                }
            })
        return JsonResponse({'error': 'Identifiants invalides'}, status=401)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)

@csrf_exempt
@require_http_methods(["POST"])
def api_register(request):
    try:
        data = json.loads(request.body)
        email = data.get('email')
        password = data.get('password')
        username = data.get('username')
        full_name = data.get('full_name', '')
        
        # Inscription avec Supabase Auth
        response = supabase.auth.sign_up({
            'email': email,
            'password': password,
        })

        # Normalize and fallback if Supabase refuses or rate-limits
        normalized = normalize_supabase_auth_response(response)
        user = normalized.get('user')
        error_msg = normalized.get('error')

        if not user and (not error_msg or 'rate limit' in str(error_msg).lower() or 'email rate limit' in str(error_msg).lower()):
            fallback_user = admin_signup_fallback(email, password)
            if fallback_user:
                user = fallback_user

        if not user:
            # If Supabase did not return a user, we rely on admin fallback only (no local fallback in Supabase-only mode)
            pass

        if user:
            upsert_user_profile(getattr(user, 'id', None) or getattr(user, 'id', ''), username=username, full_name=full_name, role='user')
            return JsonResponse({
                'message': 'Compte créé avec succès',
                'user': {
                    'id': getattr(user, 'id', None),
                    'email': getattr(user, 'email', None),
                    'username': username,
                }
            })
        return JsonResponse({'error': 'Erreur lors de l\'inscription'}, status=400)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)

@login_required_api
def api_fiches_list(request):
    try:
        log_file = open('api_fiches_debug.log', 'a')
        log_file.write(f"\n[API_FICHES] Appel reçu\n")
        
        query = supabase.table('fiches').select(
            'id,reference,name,description,category,manufacturer,version,file_name,file_type,file_url,created_at,author_id'
        )
        log_file.write(f"[API_FICHES] Query constructée\n")
        
        # Filtres
        search = request.GET.get('search', '')
        category = request.GET.get('category', '')
        tag = request.GET.get('tag', '')
        log_file.write(f"[API_FICHES] Params search={search!r}, category={category!r}, tag={tag!r}\n")
        
        if search:
            query = query.or_(f"reference.ilike.%{search}%,name.ilike.%{search}%,description.ilike.%{search}%,manufacturer.ilike.%{search}%")
        if category:
            query = query.eq('category', category)
        if tag:
            tag_response = supabase.table('tags').select('id').eq('name', tag).execute()
            if not tag_response.data:
                return JsonResponse({'results': [], 'total': 0, 'page': int(request.GET.get('page', 1))})
            tag_id = tag_response.data[0]['id']
            fiche_tags_response = supabase.table('fiche_tags').select('fiche_id').eq('tag_id', tag_id).execute()
            fiche_ids = [item['fiche_id'] for item in fiche_tags_response.data] if fiche_tags_response.data else []
            if not fiche_ids:
                return JsonResponse({'results': [], 'total': 0, 'page': int(request.GET.get('page', 1))})
            query = query.in_('id', fiche_ids)
        
        # Pagination
        page = int(request.GET.get('page', 1))
        per_page = 20
        offset = (page - 1) * per_page
        
        log_file.write(f"[API_FICHES] Avant execute - offset={offset}, per_page={per_page}\n")
        response = query.range(offset, offset + per_page - 1).order('created_at', desc=True).execute()
        log_file.write(f"[API_FICHES] Response reçue: status={getattr(response, 'status_code', None)}, rows={len(response.data) if response.data else 0}\n")

        if getattr(response, 'status_code', 200) >= 400:
            log_file.write(f"[API_FICHES] Query error: {getattr(response, 'error', None)}\n")
            log_file.close()
            return JsonResponse({'error': 'Erreur de requête Supabase', 'details': getattr(response, 'error', None)}, status=500)

        fiche_ids = [fiche.get('id') for fiche in response.data if fiche.get('id')]
        author_ids = [fiche.get('author_id') for fiche in response.data if fiche.get('author_id')]
        author_map = {}
        if author_ids:
            profiles_resp = supabase.table('profiles').select('id,username,full_name').in_('id', list(set(author_ids))).execute()
            if profiles_resp.data:
                author_map = {
                    profile['id']: {
                        'username': profile.get('username', ''),
                        'full_name': profile.get('full_name', '')
                    }
                    for profile in profiles_resp.data
                }

        tag_map = {}
        fiche_tags_map = {fiche_id: [] for fiche_id in fiche_ids}
        if fiche_ids:
            tags_resp = supabase.table('fiche_tags').select('fiche_id,tag_id').in_('fiche_id', fiche_ids).execute()
            tag_ids = [item['tag_id'] for item in tags_resp.data] if tags_resp.data else []
            tags_data = []
            if tag_ids:
                tags_data = supabase.table('tags').select('id,name,color').in_('id', list(set(tag_ids))).execute().data or []
                tag_map = {tag['id']: {'id': tag['id'], 'name': tag.get('name', ''), 'color': tag.get('color', '#293462')} for tag in tags_data}
            for item in tags_resp.data or []:
                fiche_id = item.get('fiche_id')
                tag_id = item.get('tag_id')
                if fiche_id and tag_id and tag_id in tag_map:
                    fiche_tags_map.setdefault(fiche_id, []).append(tag_map[tag_id])

        version_counts = {}
        if fiche_ids:
            versions_resp = supabase.table('versions').select('fiche_id').in_('fiche_id', fiche_ids).execute()
            for item in versions_resp.data or []:
                fid = item.get('fiche_id')
                if fid:
                    version_counts[fid] = version_counts.get(fid, 0) + 1

        results = []
        for fiche in response.data:
            log_file.write(f"[API_FICHES] Traitement fiche: {fiche.get('id')}\n")
            author_profile = author_map.get(fiche.get('author_id')) if fiche.get('author_id') else None
            author_name = None
            if author_profile:
                author_name = (author_profile.get('full_name') or '').strip() or author_profile.get('username')
            # Fallback to a friendly label when author is missing
            if not author_name:
                author_name = 'Inconnu'
            fiche_tags = fiche_tags_map.get(fiche.get('id'), [])
            versions_count = version_counts.get(fiche.get('id'), 0)

            category_value = fiche.get('category', '')
            category_display = {
                'ELECTRONIQUE': 'Électronique',
                'MECANIQUE': 'Mécanique',
                'LOGICIEL': 'Logiciel',
                'MATERIAUX': 'Matériaux',
                'ELECTRIQUE': 'Électrique',
                'AUTRE': 'Autre',
            }.get(category_value, category_value or 'Non catégorisé')

            results.append({
                'id': fiche['id'],
                'reference': fiche.get('reference', ''),
                'name': fiche.get('name', ''),
                'description': fiche.get('description', ''),
                'category': category_value,
                'category_display': category_display,
                'manufacturer': fiche.get('manufacturer', ''),
                'version': fiche.get('version', '1.0'),
                'tags': fiche_tags,
                'author': author_name,
                'author_display': author_name,
                'created_at': fiche.get('created_at', ''),
                'versions_count': versions_count,
                'file_name': fiche.get('file_name', ''),
                'file_type': fiche.get('file_type', ''),
                'file_url': fiche.get('file_url', ''),
                'view_url': f"/view/{fiche['id']}/",
            })

        log_file.write(f"[API_FICHES] {len(results)} fiches retournées\n")
        log_file.close()

        return JsonResponse({
            'results': results,
            'total': len(response.data),
            'page': page,
        })
    except Exception as e:
        log_file = open('api_fiches_debug.log', 'a')
        log_file.write(f"[API_FICHES] Exception: {str(e)}\n")
        import traceback
        log_file.write(f"[API_FICHES] Traceback: {traceback.format_exc()}\n")
        log_file.close()
        return JsonResponse({'error': str(e)}, status=500)

@login_required_api
def api_fiche_detail(request, fiche_id):
    try:
        response = supabase.table('fiches').select('*').eq('id', fiche_id).execute()
        if not response.data:
            return JsonResponse({'error': 'Fiche non trouvée'}, status=404)
        
        fiche = response.data[0]
        profile = get_user_profile(fiche.get('author_id')) if fiche.get('author_id') else None
        author_display = get_profile_display_name(profile)
        
        # Tags
        tags_response = supabase.table('fiche_tags').select('tags(*)').eq('fiche_id', fiche_id).execute()
        fiche['tags'] = [t['tags'] for t in tags_response.data] if tags_response.data else []
        
        # Versions
        versions_response = supabase.table('versions').select('*').eq('fiche_id', fiche_id).order('created_at', desc=True).execute()
        versions = versions_response.data if versions_response.data else []
        if versions:
            version_author_ids = [version.get('author_id') for version in versions if version.get('author_id')]
            version_profiles = {}
            if version_author_ids:
                version_profiles_resp = supabase.table('profiles').select('id,username,full_name').in_('id', list(set(version_author_ids))).execute()
                version_profiles = {
                    profile['id']: profile for profile in (version_profiles_resp.data or []) if profile.get('id')
                }
            for version in versions:
                version_profile = version_profiles.get(version.get('author_id'))
                version['author_display'] = get_profile_display_name(version_profile)
        fiche['versions'] = versions
        
        # Preview
        preview = fiche.get('file_preview', {})
        if isinstance(preview, str):
            try:
                preview = json.loads(preview)
            except:
                preview = {}
        
        return JsonResponse({
            'id': fiche['id'],
            'reference': fiche['reference'],
            'name': fiche['name'],
            'description': fiche.get('description', ''),
            'category': fiche.get('category', ''),
            'category_display': fiche.get('category', 'Non catégorisé'),
            'manufacturer': fiche.get('manufacturer', ''),
            'version': fiche.get('version', '1.0'),
            'file_url': fiche.get('file_url', ''),
            'file_name': fiche.get('file_name', ''),
            'file_size': fiche.get('file_size', 0),
            'file_type': fiche.get('file_type', ''),
            'file_preview': preview,
            'tags': fiche.get('tags', []),
            'author': author_display,
            'author_display': author_display,
            'created_at': fiche['created_at'],
            'updated_at': fiche.get('updated_at', fiche['created_at']),
            'versions': fiche.get('versions', []),
        })
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required_api
@csrf_exempt
@require_http_methods(["POST"])
def api_fiche_create(request):
    import traceback
    log_file = open('api_import_debug.log', 'a')
    try:
        log_file.write(f"\n[API_FICHE_CREATE] Appel reçu - user_id={getattr(request, 'user_id', None)}\n")
        log_file.write(f"POST keys: {list(request.POST.keys())}\n")
        log_file.write(f"FILES keys: {list(request.FILES.keys())}\n")
        log_file.flush()

        file = request.FILES.get('file')
        valid, error = validate_pdf_upload(file)
        if not valid:
            log_file.write(f"[API_FICHE_CREATE] Validation failed: {error}\n")
            log_file.close()
            return JsonResponse({'error': error}, status=400)

        name = request.POST.get('name')
        if not name:
            return JsonResponse({'error': 'Le nom est requis'}, status=400)
        
        reference = (request.POST.get('reference') or '').strip()
        if not reference:
            reference = f"FT-{uuid.uuid4().hex[:8]}"
        else:
            existing_reference = supabase.table('fiches').select('id').eq('reference', reference).execute()
            if existing_reference.data:
                log_file.write(f"[API_FICHE_CREATE] Duplicate reference detected: {reference}\n")
                log_file.close()
                return JsonResponse({
                    'error': 'La référence existe déjà. Choisissez une autre référence ou laissez ce champ vide pour générer automatiquement.'
                }, status=409)
        
        # Upload du fichier vers Supabase Storage
        file_data = file.read()
        if len(file_data) > MAX_FILE_SIZE:
            return JsonResponse({'error': 'Le fichier doit faire 50 MB maximum'}, status=400)
        upload_result = supabase_upload_file(file_data, file.name)
        
        if not upload_result['success']:
            return JsonResponse({'error': upload_result.get('error', 'Erreur d\'upload')}, status=500)
        
        if upload_result.get('warning'):
            log_file.write(f"[API_FICHE_CREATE] Warning: {upload_result.get('warning')}\n")
            log_file.write(f"[API_FICHE_CREATE] upload_error: {upload_result.get('upload_error')}\n")
            log_file.flush()
        
        # Déterminer le type MIME
        file_type = None
        if magic is not None:
            try:
                file_type = magic.Magic(mime=True).from_buffer(file_data)
            except Exception:
                file_type = None
        if not file_type:
            file_type = guess_mime_from_extension(file.name)
        
        # Création de la fiche dans Supabase
        fiche_data = {
            'reference': reference,
            'name': name,
            'description': request.POST.get('description', ''),
            'category': request.POST.get('category', ''),
            'manufacturer': request.POST.get('manufacturer', ''),
            'version': request.POST.get('version', '1.0'),
            'file_url': upload_result.get('url'),
            'file_name': upload_result.get('name'),
            'file_size': len(file_data),
            'file_type': file_type,
            'author_id': request.user_id,
        }
        
        # Génération de l'aperçu
        preview = generate_file_preview(file_data, file_type)
        if preview:
            fiche_data['file_preview'] = json.dumps(preview)
        
        log_file.write(f"[API_FICHE_CREATE] fiche_data: {json.dumps(fiche_data, default=str)}\n")
        log_file.flush()
        response = supabase.table('fiches').insert(fiche_data).execute()
        
        fiche = None
        if response.data:
            fiche = response.data[0]
        elif getattr(response, 'status_code', None) in (200, 201, 204):
            log_file.write(f"[API_FICHE_CREATE] Insert succeeded without body, fetching fiche by reference {reference}\n")
            lookup = supabase.table('fiches').select('*').eq('reference', reference).execute()
            if lookup.data:
                fiche = lookup.data[0]
            else:
                log_file.write("[API_FICHE_CREATE] Aucun résultat trouvé après insert sans body\n")
        
        if not fiche:
            status = getattr(response, 'status_code', None)
            err = getattr(response, 'error', None)
            log_file.write(f"[API_FICHE_CREATE] Supabase insert returned no data - status={status}, error={err}, repr={response}\n")
            # If failure due to foreign key on author_id, retry without author to avoid failing the whole import
            if status == 409 and err and ('foreign key' in str(err).lower() or 'fiches_author_id_fkey' in str(err).lower() or 'violates foreign key' in str(err).lower()):
                log_file.write('[API_FICHE_CREATE] FK violation detected on author_id, retrying insert without author_id\n')
                fiche_data_no_author = dict(fiche_data)
                fiche_data_no_author.pop('author_id', None)
                retry_resp = supabase.table('fiches').insert(fiche_data_no_author).execute()
                if retry_resp.data:
                    fiche = retry_resp.data[0]
                elif getattr(retry_resp, 'status_code', None) in (200,201,204):
                    lookup = supabase.table('fiches').select('*').eq('reference', reference).execute()
                    if lookup.data:
                        fiche = lookup.data[0]
                if fiche:
                    log_file.write('[API_FICHE_CREATE] Insert succeeded without author_id after FK error\n')
            if not fiche:
                if upload_result.get('path'):
                    log_file.write(f"[API_FICHE_CREATE] Deleting uploaded storage file due to create failure: {upload_result['path']}\n")
                    supabase_delete_file(upload_result['path'])
                log_file.close()
                if status == 409 and err and 'reference' in str(err).lower():
                    return JsonResponse({
                        'error': 'La référence existe déjà. Choisissez une autre référence ou laissez ce champ vide pour générer automatiquement.'
                    }, status=409)
                return JsonResponse({'error': 'Erreur lors de la création'}, status=500)
        
        # Tags
        tags_input = request.POST.get('tags', '')
        if tags_input:
            tag_names = [t.strip() for t in tags_input.split(',') if t.strip()]
            for tag_name in tag_names:
                # Créer ou récupérer le tag
                tag_response = supabase.table('tags').select('id').eq('name', tag_name).execute()
                if tag_response.data:
                    tag_id = tag_response.data[0]['id']
                else:
                    tag_insert = supabase.table('tags').insert({'name': tag_name}).execute()
                    tag_id = tag_insert.data[0]['id'] if tag_insert.data else None
                
                if tag_id:
                    supabase.table('fiche_tags').insert({
                        'fiche_id': fiche['id'],
                        'tag_id': tag_id
                    }).execute()
        
        # Notification
        create_notification_supabase(
            request.user_id,
            'IMPORT',
            f"Fiche importée : {fiche['reference']} - {fiche['name']}",
            fiche['id']
        )
        
        log_file.write(f"[API_FICHE_CREATE] Fiche créée: {fiche.get('id')}\n")
        log_file.close()

        return JsonResponse({
            'message': 'Fiche créée avec succès',
            'fiche': {
                'id': fiche['id'],
                'reference': fiche['reference'],
                'name': fiche['name'],
                'view_url': f"/view/{fiche['id']}/",
            }
        }, status=201)
    except Exception as e:
        log_file.write(f"[API_FICHE_CREATE] Exception: {str(e)}\n")
        log_file.write(traceback.format_exc())
        log_file.close()
        return JsonResponse({'error': str(e)}, status=500)

@login_required_api
@csrf_exempt
@require_http_methods(["PUT"])
def api_fiche_update(request, fiche_id):
    try:
        open('api_fiche_update.log', 'a', encoding='utf-8').write(f"[API_FICHE_UPDATE] called fiche_id={fiche_id} user_id={getattr(request,'user_id',None)}\n")
        data = json.loads(request.body)
        
        update_data = {}
        if 'name' in data:
            update_data['name'] = data['name']
        if 'description' in data:
            update_data['description'] = data['description']
        if 'category' in data:
            update_data['category'] = data['category']
        if 'manufacturer' in data:
            update_data['manufacturer'] = data['manufacturer']
        update_data['updated_at'] = datetime.now().isoformat()
        
        if update_data:
            supabase.table('fiches').update(update_data).eq('id', fiche_id).execute()
        
        # Mise à jour des tags
        if 'tags' in data:
            # Supprimer les anciens tags
            supabase.table('fiche_tags').delete().eq('fiche_id', fiche_id).execute()
            
            for tag_name in data['tags']:
                if tag_name.strip():
                    tag_response = supabase.table('tags').select('id').eq('name', tag_name.strip()).execute()
                    if tag_response.data:
                        tag_id = tag_response.data[0]['id']
                    else:
                        tag_insert = supabase.table('tags').insert({'name': tag_name.strip()}).execute()
                        tag_id = tag_insert.data[0]['id'] if tag_insert.data else None
                    
                    if tag_id:
                        supabase.table('fiche_tags').insert({
                            'fiche_id': fiche_id,
                            'tag_id': tag_id
                        }).execute()
        
        create_notification_supabase(
            request.user_id,
            'UPDATE',
            f"Fiche mise à jour",
            fiche_id
        )
        
        return JsonResponse({'message': 'Fiche mise à jour avec succès'})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)

@login_required_api
@require_http_methods(["DELETE"])
def api_fiche_delete(request, fiche_id):
    try:
        # Récupérer la fiche pour supprimer le fichier
        response = supabase.table('fiches').select('file_url, file_name').eq('id', fiche_id).execute()
        if response.data:
            fiche = response.data[0]
            if fiche.get('file_url'):
                file_path = extract_storage_path(fiche['file_url'])
                supabase_delete_file(file_path)
        
        # Supprimer les tags associés
        supabase.table('fiche_tags').delete().eq('fiche_id', fiche_id).execute()
        
        # Supprimer les versions
        versions = supabase.table('versions').select('file_url').eq('fiche_id', fiche_id).execute()
        for v in versions.data:
            if v.get('file_url'):
                file_path = extract_storage_path(v['file_url'])
                supabase_delete_file(file_path)
        supabase.table('versions').delete().eq('fiche_id', fiche_id).execute()
        
        # Supprimer la fiche
        supabase.table('fiches').delete().eq('id', fiche_id).execute()
        
        create_notification_supabase(
            request.user_id,
            'DELETE',
            f"Fiche supprimée",
            fiche_id
        )
        
        return JsonResponse({'message': 'Fiche supprimée avec succès'})
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)

@login_required_api
@csrf_exempt
@require_http_methods(["POST"])
def api_fiche_version(request, fiche_id):
    try:
        file = request.FILES.get('file')
        valid, error = validate_pdf_upload(file)
        if not valid:
            return JsonResponse({'error': error}, status=400)
        
        # Récupérer la fiche
        response = supabase.table('fiches').select('version').eq('id', fiche_id).execute()
        if not response.data:
            return JsonResponse({'error': 'Fiche non trouvée'}, status=404)
        
        current_version = response.data[0].get('version', '1.0')
        
        # Incrémenter la version
        version_parts = current_version.split('.')
        if len(version_parts) == 2:
            new_version = f"{version_parts[0]}.{int(version_parts[1]) + 1}"
        else:
            new_version = "1.1"
        
        # Upload du fichier
        file_data = file.read()
        upload_result = supabase_upload_file(file_data, file.name)
        
        if not upload_result['success']:
            return JsonResponse({'error': upload_result.get('error', 'Erreur d\'upload')}, status=500)
        
        # Créer la version
        version_data = {
            'fiche_id': fiche_id,
            'version': new_version,
            'file_url': upload_result['url'],
            'file_name': upload_result['name'],
            'comment': request.POST.get('comment', ''),
            'author_id': request.user_id,
        }
        supabase.table('versions').insert(version_data).execute()
        
        # Mettre à jour la fiche
        supabase.table('fiches').update({
            'version': new_version,
            'file_url': upload_result['url'],
            'file_name': upload_result['name'],
            'updated_at': datetime.now().isoformat()
        }).eq('id', fiche_id).execute()
        
        create_notification_supabase(
            request.user_id,
            'VERSION',
            f"Nouvelle version {new_version}",
            fiche_id
        )
        
        return JsonResponse({
            'message': 'Version ajoutée avec succès',
            'version': new_version
        })
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)

def get_mime_type(filename):
    """Détermine le MIME type d'un fichier"""
    mime_types = {
        '.pdf': 'application/pdf',
        '.doc': 'application/msword',
        '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        '.xls': 'application/vnd.ms-excel',
        '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        '.ppt': 'application/vnd.ms-powerpoint',
        '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        '.txt': 'text/plain',
        '.csv': 'text/csv',
        '.png': 'image/png',
        '.jpg': 'image/jpeg',
        '.jpeg': 'image/jpeg',
        '.gif': 'image/gif',
        '.zip': 'application/zip',
    }
    ext = Path(filename).suffix.lower()
    return mime_types.get(ext, 'application/octet-stream')

@login_required_api
def api_fiche_download(request, fiche_id):
    try:
        response = supabase.table('fiches').select('file_url, file_name, file_type').eq('id', fiche_id).execute()
        if not response.data or not response.data[0].get('file_url'):
            return JsonResponse({'error': 'Fichier non trouvé'}, status=404)
        
        fiche = response.data[0]
        file_url = fiche['file_url'].replace('/media/uploads/uploads/', '/media/uploads/')
        file_name = fiche.get('file_name', 'document')
        file_type = fiche.get('file_type') or get_mime_type(file_name)
        
        if file_url.startswith('/media/'):
            local_path = BASE_DIR / file_url.lstrip('/')
            if local_path.exists():
                response = FileResponse(local_path.open('rb'), content_type=file_type)
                response['Content-Disposition'] = f'attachment; filename="{file_name}"'
                return response
        else:
            fallback_path = find_local_fallback_path(file_url, file_name=file_name)
            if fallback_path:
                response = FileResponse(fallback_path.open('rb'), content_type=file_type)
                response['Content-Disposition'] = f'attachment; filename="{file_name}"'
                return response

        file_content, error, remote_content_type = download_remote_file(file_url)
        if file_content is None:
            return JsonResponse({'error': 'Erreur de téléchargement', 'details': error}, status=502)
        if not file_type and remote_content_type:
            file_type = remote_content_type
        response = HttpResponse(file_content, content_type=file_type)
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required_api
def api_version_download(request, fiche_id, version_id):
    try:
        response = supabase.table('versions').select('file_url, file_name, id').eq('id', version_id).eq('fiche_id', fiche_id).execute()
        if not response.data or not response.data[0].get('file_url'):
            return JsonResponse({'error': 'Version non trouvée'}, status=404)
        
        version = response.data[0]
        file_url = version.get('file_url', '').replace('/media/uploads/uploads/', '/media/uploads/')
        file_name = version.get('file_name', 'version')
        file_type = get_mime_type(file_name)
        
        if file_url.startswith('/media/'):
            local_path = BASE_DIR / file_url.lstrip('/')
            if local_path.exists():
                response = FileResponse(local_path.open('rb'), content_type=file_type)
                response['Content-Disposition'] = f'attachment; filename="{file_name}"'
                return response
            # Fallback vers Supabase si fichier local manquant

        file_content, error, remote_content_type = download_remote_file(file_url)
        if file_content is None:
            return JsonResponse({'error': 'Erreur de téléchargement', 'details': error}, status=502)
        if not file_type and remote_content_type:
            file_type = remote_content_type
        response = HttpResponse(file_content, content_type=file_type)
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response
        
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

def api_fiche_view(request, fiche_id):
    try:
        # Auth optional: try to populate request.user_id if token/session available,
        # but allow unauthenticated access so public files can be viewed in an iframe.
        try:
            token = request.headers.get('Authorization', '').replace('Bearer ', '')
            if token:
                user = supabase.auth.get_user(token)
                if user:
                    request.user_id = user.user.id
                    request.user_email = user.user.email
        except Exception:
            pass
        if not getattr(request, 'user_id', None):
            user_id = request.session.get('user_id')
            if user_id:
                request.user_id = user_id

        response = supabase.table('fiches').select('id,name,reference,file_url,file_type,file_preview').eq('id', fiche_id).execute()
        if not response.data:
            return JsonResponse({'error': 'Fiche non trouvée'}, status=404)
        
        fiche = response.data[0]
        preview = fiche.get('file_preview', {})
        if isinstance(preview, str):
            try:
                preview = json.loads(preview)
            except:
                preview = {}
        
        # If the stored file_url is a public HTTP(S) URL (Supabase public storage),
        # return it directly so the browser can load the PDF/image without hitting
        # the authenticated serve endpoint (avoids iframe auth issues).
        stored_url = fiche.get('file_url', '')
        if isinstance(stored_url, str) and stored_url.startswith('http'):
            file_url = stored_url
        else:
            file_url = f"/api/fiches/{fiche_id}/serve/"

        return JsonResponse({
            'id': fiche['id'],
            'name': fiche['name'],
            'reference': fiche['reference'],
            'file_type': fiche.get('file_type', ''),
            'file_url': file_url,
            'preview': preview,
        })
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required_api
def api_fiche_serve(request, fiche_id):
    """Sert le fichier directement pour visualisation"""
    try:
        fiche_response = supabase.table('fiches').select('file_url,file_name,file_type').eq('id', fiche_id).execute()
        if not fiche_response.data or not fiche_response.data[0].get('file_url'):
            return JsonResponse({'error': 'Fichier non trouvé'}, status=404)
        
        fiche = fiche_response.data[0]
        file_url = fiche.get('file_url', '')
        file_name = fiche.get('file_name', 'document')
        file_type = fiche.get('file_type') or get_mime_type(file_name)
        file_url = file_url.replace('/media/uploads/uploads/', '/media/uploads/')

        if file_url.startswith('/media/'):
            local_path = BASE_DIR / file_url.lstrip('/')
            if local_path.exists():
                with open(local_path, 'rb') as f:
                    file_content = f.read()
                response = HttpResponse(file_content, content_type=file_type)
                response['X-Frame-Options'] = 'SAMEORIGIN'
                response['Content-Disposition'] = f'inline; filename="{file_name}"'
                return response
        else:
            fallback_path = find_local_fallback_path(file_url, file_name=file_name)
            if fallback_path:
                with open(fallback_path, 'rb') as f:
                    file_content = f.read()
                response = HttpResponse(file_content, content_type=file_type)
                response['X-Frame-Options'] = 'SAMEORIGIN'
                response['Content-Disposition'] = f'inline; filename="{file_name}"'
                return response

        if file_url.startswith('http'):
            file_content, error, remote_content_type = download_remote_file(file_url)
            if file_content is None:
                return JsonResponse({'error': 'Erreur de téléchargement du fichier', 'details': error}, status=502)
            if not file_type and remote_content_type:
                file_type = remote_content_type
            response = HttpResponse(file_content, content_type=file_type)
            response['X-Frame-Options'] = 'SAMEORIGIN'
            response['Content-Disposition'] = f'inline; filename="{file_name}"'
            return response

        return JsonResponse({'error': 'Fichier non trouvé'}, status=404)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({'error': str(e)}, status=500)

@login_required_api
def api_notifications(request):
    try:
        response = supabase.table('notifications').select('*').eq('recipient_id', request.user_id).order('created_at', desc=True).limit(50).execute()
        return JsonResponse(response.data if response.data else [], safe=False)
    except:
        return JsonResponse([], safe=False)

@login_required_api
@require_http_methods(["PUT"])
def api_notifications_mark_read(request, notification_id):
    try:
        supabase.table('notifications').update({'is_read': True}).eq('id', notification_id).eq('recipient_id', request.user_id).execute()
        return JsonResponse({'message': 'Notification marquée comme lue'})
    except:
        return JsonResponse({'error': 'Erreur'}, status=500)

@login_required_api
@require_http_methods(["PUT"])
def api_notifications_mark_all_read(request):
    try:
        supabase.table('notifications').update({'is_read': True}).eq('recipient_id', request.user_id).eq('is_read', False).execute()
        return JsonResponse({'message': 'Toutes les notifications marquées comme lues'})
    except:
        return JsonResponse({'error': 'Erreur'}, status=500)

@login_required_api
def api_stats(request):
    try:
        total = supabase.table('fiches').select('id', count='exact').execute()
        categories = supabase.table('fiches').select('category').execute()
        tags = supabase.table('tags').select('id', count='exact').execute()
        fiche_tags = supabase.table('fiche_tags').select('tag_id').execute()
        
        # Compter les versions
        versions = supabase.table('versions').select('id', count='exact').execute()
        
        # Catégories distinctes
        cat_set = set()
        if categories.data:
            for c in categories.data:
                if c.get('category'):
                    cat_set.add(c['category'])
        
        # Tags utilisés distinctement
        tag_ids = set()
        if fiche_tags.data:
            for row in fiche_tags.data:
                if row.get('tag_id'):
                    tag_ids.add(row['tag_id'])
        
        return JsonResponse({
            'total_fiches': int(total.count) if hasattr(total, 'count') and total.count is not None else 0,
            'categories_count': len(cat_set),
            'tags_count': len(tag_ids),
            'versions_count': int(versions.count) if hasattr(versions, 'count') and versions.count is not None else 0,
        })
    except Exception as e:
        return JsonResponse({'total_fiches': 0, 'categories_count': 0, 'tags_count': 0, 'versions_count': 0})

@login_required_api
def api_tags(request):
    try:
        response = supabase.table('tags').select('*').execute()
        return JsonResponse(response.data if response.data else [], safe=False)
    except:
        return JsonResponse([], safe=False)

# =================================================================================
# UTILITAIRES
# =================================================================================

MAX_FILE_SIZE = 50 * 1024 * 1024


def is_pdf_file(filename, file_type=None):
    if file_type:
        if file_type.lower() == 'application/pdf':
            return True
        if file_type.lower().startswith('application/pdf'):
            return True
    ext = Path(filename).suffix.lower()
    return ext == '.pdf'


def is_supported_upload_file(filename, file_type=None):
    if file_type:
        normalized_type = file_type.lower()
        if normalized_type == 'application/pdf' or normalized_type.startswith('application/pdf'):
            return True
        if normalized_type.startswith('image/'):
            return True
        if 'word' in normalized_type or 'document' in normalized_type or 'officedocument' in normalized_type:
            return True
        if 'excel' in normalized_type or 'spreadsheet' in normalized_type:
            return True
        if 'text' in normalized_type or normalized_type in {'application/plain', 'text/plain'}:
            return True
    ext = Path(filename).suffix.lower()
    return ext in {'.pdf', '.png', '.jpg', '.jpeg', '.gif', '.doc', '.docx', '.xls', '.xlsx', '.txt', '.csv'}


def validate_pdf_upload(file):
    if not file:
        return False, 'Aucun fichier fourni'
    size = getattr(file, 'size', None)
    if size is not None and size > MAX_FILE_SIZE:
        return False, 'Le fichier doit faire 50 MB maximum'
    filename = getattr(file, 'name', '')
    file_type = getattr(file, 'content_type', '')
    if not is_supported_upload_file(filename, file_type):
        return False, 'Format non pris en charge. Utilisez PDF, image, Word, Excel ou texte.'
    return True, None


def guess_mime_from_extension(filename):
    """Détecte un type MIME simple à partir de l’extension."""
    ext = Path(filename).suffix.lower()
    mapping = {
        '.pdf': 'application/pdf',
        '.png': 'image/png',
        '.jpg': 'image/jpeg',
        '.jpeg': 'image/jpeg',
        '.gif': 'image/gif',
        '.doc': 'application/msword',
        '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        '.xls': 'application/vnd.ms-excel',
        '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        '.txt': 'text/plain',
    }
    return mapping.get(ext, 'application/octet-stream')


def generate_file_preview(file_data, file_type):
    """Génère un aperçu du fichier"""
    preview = {
        'type': file_type,
        'content': None,
        'pages': None,
        'text': None,
        'images': [],
        'sheets': None,
        'error': None
    }
    
    try:
        # PDF
        if 'pdf' in file_type:
            from io import BytesIO
            pdf_file = BytesIO(file_data)
            reader = PyPDF2.PdfReader(pdf_file)
            preview['pages'] = len(reader.pages)
            preview['text'] = ''
            for page in reader.pages[:2]:
                text = page.extract_text()
                if text:
                    preview['text'] += text[:1000] + '...'
                    break
        
        # Images
        elif file_type.startswith('image/'):
            preview['images'] = [{'width': None, 'height': None, 'format': file_type}]
        
        # Word
        elif 'word' in file_type or 'document' in file_type:
            if Document is not None:
                from io import BytesIO
                doc = Document(BytesIO(file_data))
                preview['text'] = ''
                for para in doc.paragraphs[:15]:
                    if para.text.strip():
                        preview['text'] += para.text[:200] + ' '
                preview['text'] = preview['text'][:1000] + '...' if len(preview['text']) > 1000 else preview['text']
            else:
                preview['text'] = 'Aperçu Word indisponible dans cet environnement.'
        
        # Excel
        elif 'excel' in file_type or 'spreadsheet' in file_type:
            from io import BytesIO
            wb = openpyxl.load_workbook(BytesIO(file_data), data_only=True)
            preview['sheets'] = len(wb.sheetnames)
            preview['sheet_names'] = wb.sheetnames[:5]
        
        # Texte
        elif 'text' in file_type:
            preview['text'] = file_data[:2000].decode('utf-8', errors='ignore')
        
    except Exception as e:
        preview['error'] = str(e)
    
    return preview

# =================================================================================
# TEMPLATES HTML
# =================================================================================

TEMPLATE_FILES = {
    'base.html': '''
<!DOCTYPE html>
<html lang="fr">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <meta http-equiv="Cache-Control" content="no-cache, no-store, must-revalidate">
    <meta http-equiv="Pragma" content="no-cache">
    <meta http-equiv="Expires" content="0">
    <title>{% block title %}Quantum Technology - Gestion de Fiches{% endblock %}</title>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&display=swap" rel="stylesheet">
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
    <style>
        :root {
            --primary: #D61C4E;
            --primary-dark: #B81842;
            --secondary: #293462;
            --secondary-light: #3A4A7A;
            --accent: #293462;
            --success: #D61C4E;
            --warning: #293462;
            --danger: #A31B45;
            --shadow-sm: 0 1px 2px rgba(0,0,0,0.08);
            --shadow: 0 16px 40px rgba(0,0,0,0.18);
            --shadow-lg: 0 24px 70px rgba(0,0,0,0.24);
            --radius: 16px;
            --radius-lg: 28px;
            --transition: all 0.3s cubic-bezier(0.4, 0, 0.2, 1);
            --surface: rgba(255,255,255,0.96);
            --surface-strong: rgba(255,255,255,1);
            --text: #1F2937;
            --text-muted: #475569;
        }
        * { margin: 0; padding: 0; box-sizing: border-box; }
        body {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
            min-height: 100vh;
            background: radial-gradient(circle at 15% 15%, rgba(255,255,255,0.04), transparent 18%),
                        radial-gradient(circle at 85% 10%, rgba(214,28,78,0.08), transparent 14%),
                        linear-gradient(135deg, #0b122d 0%, #14244b 40%, #192f5f 100%);
            color: var(--text);
            line-height: 1.6;
            position: relative;
            overflow-x: hidden;
        }
        body::before {
            content: '';
            position: fixed;
            inset: 0;
            background: radial-gradient(circle at 50% 40%, rgba(255,255,255,0.1), transparent 14%);
            pointer-events: none;
            z-index: 0;
        }
        * { box-sizing: border-box; }
        
        /* Logo dans la navbar */
        .nav-brand {
            display: flex;
            align-items: center;
            gap: 0.75rem;
        }
        .nav-brand .logo-icon {
            width: 80px;
            height: 80px;
            background: #ffffff;
            border-radius: 20px;
            display: flex;
            align-items: center;
            justify-content: center;
            box-shadow: 0 14px 30px rgba(0,0,0,0.12);
            overflow: hidden;
        }
        .nav-brand .logo-icon img {
            width: 100%;
            height: 100%;
            object-fit: contain;
            display: block;
        }
        .nav-brand .logo-text {
            display: flex;
            flex-direction: column;
            line-height: 1.2;
        }
        .nav-brand .logo-text .brand {
            font-weight: 800;
            font-size: 1.1rem;
            color: var(--secondary);
        }
        .nav-brand .logo-text .sub {
            font-size: 0.6rem;
            color: var(--gray-500);
            font-weight: 600;
            letter-spacing: 1px;
            text-transform: uppercase;
        }
        
        /* Navbar */
        .navbar { background: rgba(255,255,255,0.88); border-bottom: 1px solid rgba(214,28,78,0.18); padding: 0 2rem; height: 68px; position: sticky; top: 0; z-index: 1000; box-shadow: var(--shadow-sm); backdrop-filter: blur(18px); }
        .nav-container { max-width: 1440px; margin: 0 auto; height: 100%; display: flex; align-items: center; justify-content: space-between; gap: 1.5rem; }
        .nav-search { flex: 1; max-width: 500px; }
        .nav-search form { display: flex; align-items: center; background: rgba(255,255,255,0.92); border-radius: var(--radius); padding: 0.25rem; transition: var(--transition); }
        .nav-search form:focus-within { background: white; box-shadow: 0 0 0 3px rgba(214, 28, 78, 0.15); }
        .nav-search input { border: none; background: transparent; padding: 0.5rem 1rem; flex: 1; font-size: 0.9rem; outline: none; color: var(--gray-800); }
        .nav-search button { background: transparent; border: none; padding: 0.5rem 1rem; color: var(--gray-400); cursor: pointer; transition: var(--transition); }
        .nav-search button:hover { color: var(--primary); }
        .nav-right { display: flex; align-items: center; gap: 1rem; }
        .nav-btn { background: transparent; border: none; padding: 0.5rem; font-size: 1.2rem; color: var(--gray-500); cursor: pointer; position: relative; border-radius: var(--radius); transition: var(--transition); }
        .nav-btn:hover { background: var(--gray-100); color: var(--primary); }
        .notif-badge { position: absolute; top: 0; right: 0; background: var(--primary); color: white; font-size: 0.6rem; font-weight: 600; padding: 0.1rem 0.4rem; border-radius: 50%; min-width: 18px; text-align: center; transform: translate(25%, -25%); }
        .user-menu { display: flex; align-items: center; gap: 0.75rem; cursor: pointer; padding: 0.25rem 0.5rem; border-radius: var(--radius); transition: var(--transition); position: relative; }
        .user-menu:hover { background: var(--gray-100); }
        .user-avatar { width: 36px; height: 36px; border-radius: 50%; background: var(--secondary); color: white; display: flex; align-items: center; justify-content: center; font-weight: 600; font-size: 0.9rem; flex-shrink: 0; }
        .user-info { line-height: 1.3; }
        .user-name { font-weight: 500; font-size: 0.9rem; display: block; }
        .user-role { font-size: 0.7rem; color: var(--gray-400); }
        .dropdown { display: none; position: absolute; top: 100%; right: 0; margin-top: 0.5rem; background: white; border-radius: var(--radius); box-shadow: var(--shadow-lg); min-width: 200px; padding: 0.5rem; z-index: 1001; border: 1px solid var(--gray-200); }
        .dropdown.active { display: block; animation: slideDown 0.2s ease; }
        .dropdown a { display: flex; align-items: center; gap: 0.75rem; padding: 0.5rem 1rem; text-decoration: none; color: var(--gray-700); border-radius: var(--radius); transition: var(--transition); }
        .dropdown a:hover { background: var(--gray-100); color: var(--primary); }
        .dropdown a.text-danger { color: var(--danger); }
        .dropdown a.text-danger:hover { background: rgba(214,28,78,0.12); }
        .dropdown hr { margin: 0.5rem 0; border: none; border-top: 1px solid var(--gray-200); }
        @keyframes slideDown { from { opacity: 0; transform: translateY(-10px); } to { opacity: 1; transform: translateY(0); } }
        
        /* Notifications Panel */
        .notif-panel { display: none; position: fixed; top: 76px; right: 2rem; width: 380px; max-height: 450px; background: white; border-radius: var(--radius-lg); box-shadow: var(--shadow-lg); z-index: 999; overflow: hidden; border: 1px solid var(--gray-200); }
        .notif-panel.active { display: block; animation: slideDown 0.3s ease; }
        .notif-header { display: flex; justify-content: space-between; align-items: center; padding: 1rem 1.25rem; border-bottom: 1px solid var(--gray-200); }
        .notif-header h3 { font-size: 1rem; font-weight: 600; display: flex; align-items: center; gap: 0.5rem; }
        .notif-header h3 i { color: var(--primary); }
        .btn-text { background: transparent; border: none; color: var(--primary); font-size: 0.85rem; cursor: pointer; padding: 0.25rem 0.5rem; border-radius: var(--radius); transition: var(--transition); }
        .btn-text:hover { background: var(--gray-100); }
        .notif-list { overflow-y: auto; max-height: 350px; padding: 0.5rem; }
        .notif-item { display: flex; gap: 0.75rem; padding: 0.75rem; border-radius: var(--radius); transition: var(--transition); cursor: default; }
        .notif-item:hover { background: var(--gray-100); }
        .notif-item.unread { border-left: 3px solid var(--primary); background: rgba(214,28,78,0.08); }
        .notif-icon { width: 32px; height: 32px; border-radius: 50%; display: flex; align-items: center; justify-content: center; flex-shrink: 0; font-size: 0.8rem; }
        .notif-icon.IMPORT { background: rgba(214,28,78,0.16); color: var(--primary); }
        .notif-icon.UPDATE { background: rgba(41,52,98,0.18); color: var(--secondary); }
        .notif-icon.DELETE { background: rgba(214,28,78,0.16); color: var(--danger); }
        .notif-icon.VERSION { background: rgba(41,52,98,0.12); color: var(--secondary); }
        .notif-content { flex: 1; }
        .notif-content .message { font-size: 0.9rem; }
        .notif-content .time { font-size: 0.75rem; color: var(--gray-400); }
        .notif-empty { text-align: center; padding: 2rem 1rem; color: var(--gray-400); }
        .notif-empty i { font-size: 2rem; display: block; margin-bottom: 0.5rem; }
        
        /* Main */
        .main-content {
            position: relative;
            z-index: 1;
            max-width: 1440px;
            margin: 2.5rem auto 3rem;
            padding: 2rem;
            background: var(--surface);
            border-radius: 32px;
            box-shadow: var(--shadow);
            border: 1px solid rgba(255,255,255,0.55);
            backdrop-filter: blur(18px);
        }

        .hero-banner {
            display: flex;
            justify-content: space-between;
            align-items: center;
            gap: 1.5rem;
            background: linear-gradient(135deg, #0f1732 0%, #172d55 50%, #1f3a68 100%);
            border-radius: 32px;
            padding: 2.5rem;
            box-shadow: 0 24px 70px rgba(15,23,42,0.18);
            border: 1px solid rgba(255,255,255,0.08);
            margin-bottom: 2rem;
            color: white;
        }
        .hero-banner h1 {
            font-size: clamp(2rem, 2.6vw, 3.4rem);
            line-height: 1.05;
            color: white !important;
            margin-bottom: 0.75rem;
        }
        .hero-banner p {
            max-width: 680px;
            font-size: 1rem;
            color: rgba(255,255,255,0.92) !important;
            margin-bottom: 1.25rem;
        }
        .hero-content { display: flex; flex-wrap: wrap; align-items: flex-start; justify-content: space-between; gap: 2rem; }
        .hero-content > div { max-width: 720px; }
        .hero-banner .btn-primary { padding: 0.95rem 1.5rem; font-size: 0.95rem; background: transparent; border: 1px solid rgba(255,255,255,0.95); color: white !important; }
        .hero-banner .btn-primary:hover { background: rgba(255,255,255,0.14); border-color: rgba(255,255,255,0.95); color: white !important; }
        .hero-banner .btn-primary i { margin-right: 0.5rem; }

        /* Stats */
        .stats-bar { display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 1.25rem; margin-bottom: 2rem; }
        .stat-card { background: white; padding: 1.25rem; border-radius: var(--radius); box-shadow: var(--shadow-sm); transition: var(--transition); border-left: 4px solid var(--primary); }
        .stat-card:hover { box-shadow: var(--shadow); transform: translateY(-2px); }
        .stat-card .stat-header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 0.5rem; }
        .stat-card .stat-header .label { font-size: 0.8rem; color: var(--gray-500); font-weight: 500; text-transform: uppercase; letter-spacing: 0.5px; }
        .stat-card .stat-header .icon { width: 36px; height: 36px; border-radius: 50%; display: flex; align-items: center; justify-content: center; font-size: 1rem; }
        .stat-card .stat-header .icon.primary { background: rgba(214,28,78,0.14); color: var(--primary); }
        .stat-card .stat-header .icon.secondary { background: rgba(41,52,98,0.14); color: var(--secondary); }
        .stat-card .stat-header .icon.success { background: rgba(214,28,78,0.12); color: var(--primary); }
        .stat-card .stat-header .icon.accent { background: rgba(41,52,98,0.12); color: var(--accent); }
        .stat-card .stat-value { font-size: 1.75rem; font-weight: 700; color: var(--gray-900); }
        .stat-card .stat-footer { margin-top: 0.25rem; font-size: 0.75rem; color: var(--gray-400); }
        
        /* Buttons */
        .btn { display: inline-flex; align-items: center; gap: 0.5rem; padding: 0.5rem 1.25rem; border: none; border-radius: var(--radius); font-size: 0.85rem; font-weight: 500; cursor: pointer; transition: var(--transition); text-decoration: none; color: white; }
        .btn-primary { background: var(--primary); color: white; }
        .btn-primary:hover { background: var(--primary-dark); transform: translateY(-1px); box-shadow: var(--shadow); }
        .btn-secondary { background: var(--secondary); color: white; }
        .btn-secondary:hover { background: var(--secondary-light); transform: translateY(-1px); box-shadow: var(--shadow); }
        .btn-success { background: var(--primary); color: white; }
        .btn-success:hover { background: var(--primary-dark); transform: translateY(-1px); box-shadow: var(--shadow); }
        .btn-accent { background: var(--secondary); color: white; }
        .btn-accent:hover { background: var(--secondary-light); transform: translateY(-1px); box-shadow: var(--shadow); }
        .btn-outline { background: transparent; color: var(--secondary); border: 1px solid var(--secondary); }
        .btn-outline:hover { background: rgba(41,52,98,0.08); color: var(--secondary); }
        .btn-danger { background: var(--danger); color: white; }
        .btn-danger:hover { background: #9B1639; transform: translateY(-1px); box-shadow: var(--shadow); }
        .btn-sm { padding: 0.3rem 0.75rem; font-size: 0.8rem; }
        
        /* Cards Grid */
        .fiches-grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(340px, 1fr)); gap: 1.5rem; }
        .fiche-card { background: white; border-radius: var(--radius-lg); box-shadow: var(--shadow-sm); transition: var(--transition); overflow: hidden; border: 1px solid var(--gray-200); }
        .fiche-card:hover { box-shadow: var(--shadow-lg); transform: translateY(-4px); }
        .fiche-card .card-header { padding: 1.25rem 1.25rem 0.75rem; display: flex; justify-content: space-between; align-items: flex-start; gap: 1rem; }
        .fiche-card .card-header .title { flex: 1; }
        .fiche-card .card-header .title h3 { font-size: 1rem; font-weight: 600; color: var(--gray-900); margin-bottom: 0.25rem; }
        .fiche-card .card-header .title .reference { font-size: 0.8rem; color: var(--gray-400); font-family: monospace; }
        .fiche-card .version-badge { background: var(--gray-100); padding: 0.2rem 0.6rem; border-radius: 12px; font-size: 0.7rem; font-weight: 600; color: var(--gray-600); white-space: nowrap; }
        .fiche-card .card-body { padding: 0.75rem 1.25rem; }
        .fiche-card .card-body .meta { font-size: 0.8rem; color: var(--gray-500); margin-bottom: 0.5rem; }
        .fiche-card .card-body .meta i { width: 1rem; margin-right: 0.25rem; color: var(--primary); }
        .fiche-card .card-body .description { font-size: 0.85rem; color: var(--gray-600); display: -webkit-box; -webkit-line-clamp: 2; -webkit-box-orient: vertical; overflow: hidden; margin: 0.5rem 0; }
        .fiche-card .card-body .tags { display: flex; flex-wrap: wrap; gap: 0.3rem; margin: 0.5rem 0; }
        .fiche-card .tag { padding: 0.15rem 0.6rem; border-radius: 12px; font-size: 0.7rem; font-weight: 500; color: white; }
        .fiche-card .card-actions { padding: 0.75rem 1.25rem 1.25rem; display: flex; flex-wrap: wrap; gap: 0.4rem; border-top: 1px solid var(--gray-200); }
        .fiche-card .card-footer { padding: 0.5rem 1.25rem; background: var(--gray-50); display: flex; justify-content: space-between; font-size: 0.75rem; color: var(--gray-400); border-top: 1px solid var(--gray-200); }
        
        /* Empty State */
        .empty-state { text-align: center; padding: 4rem 2rem; color: var(--gray-500); grid-column: 1 / -1; }
        .empty-state i { font-size: 3rem; color: var(--gray-300); margin-bottom: 1rem; display: block; }
        .empty-state h3 { font-size: 1.1rem; color: var(--gray-700); margin-bottom: 0.5rem; }
        .empty-state p { font-size: 0.9rem; margin-bottom: 1.5rem; }
        
        /* Loading */
        .loading-spinner { display: flex; justify-content: center; align-items: center; padding: 3rem; grid-column: 1 / -1; }
        .spinner { width: 40px; height: 40px; border: 3px solid var(--gray-200); border-top-color: var(--primary); border-radius: 50%; animation: spin 0.8s linear infinite; }
        @keyframes spin { to { transform: rotate(360deg); } }
        
        /* Modal */
        .modal-overlay { display: none; position: fixed; inset: 0; background: rgba(0,0,0,0.5); z-index: 2000; align-items: center; justify-content: center; padding: 1rem; backdrop-filter: blur(4px); }
        .modal-overlay.active { display: flex; }
        .modal { background: white; border-radius: var(--radius-lg); max-width: 600px; width: 100%; max-height: 90vh; overflow-y: auto; animation: modalSlide 0.3s ease; }
        @keyframes modalSlide { from { opacity: 0; transform: translateY(20px) scale(0.95); } to { opacity: 1; transform: translateY(0) scale(1); } }
        .modal-header { padding: 1.25rem 1.5rem; border-bottom: 1px solid var(--gray-200); display: flex; justify-content: space-between; align-items: center; position: sticky; top: 0; background: white; z-index: 1; border-radius: var(--radius-lg) var(--radius-lg) 0 0; }
        .modal-header h2 { font-size: 1.1rem; font-weight: 600; display: flex; align-items: center; gap: 0.5rem; }
        .modal-header h2 i { color: var(--primary); }
        .modal-header .close-btn { background: transparent; border: none; font-size: 1.5rem; color: var(--gray-400); cursor: pointer; padding: 0.25rem; transition: var(--transition); }
        .modal-header .close-btn:hover { color: var(--gray-700); }
        .modal-body { padding: 1.5rem; }
        .modal-footer { padding: 1rem 1.5rem; border-top: 1px solid var(--gray-200); display: flex; gap: 0.75rem; justify-content: flex-end; background: var(--gray-50); border-radius: 0 0 var(--radius-lg) var(--radius-lg); }
        
        /* Form */
        .form-group { margin-bottom: 1.25rem; }
        .form-group label { display: block; font-size: 0.85rem; font-weight: 500; color: var(--gray-700); margin-bottom: 0.3rem; }
        .form-group label .required { color: var(--primary); }
        .form-group input, .form-group select, .form-group textarea { width: 100%; padding: 0.6rem 0.8rem; border: 1px solid var(--gray-200); border-radius: var(--radius); font-size: 0.9rem; transition: var(--transition); font-family: inherit; }
        .form-group input:focus, .form-group select:focus, .form-group textarea:focus { outline: none; border-color: var(--primary); box-shadow: 0 0 0 3px rgba(214, 28, 78, 0.1); }
        .form-group textarea { min-height: 80px; resize: vertical; }
        .form-row { display: grid; grid-template-columns: 1fr 1fr; gap: 1rem; }
        
        /* File Upload */
        .file-drop-zone { border: 2px dashed var(--gray-200); border-radius: var(--radius); padding: 2rem; text-align: center; cursor: pointer; transition: var(--transition); }
        .file-drop-zone:hover { border-color: var(--primary); background: rgba(214,28,78,0.08); }
        .file-drop-zone.dragover { border-color: var(--primary); background: rgba(214,28,78,0.12); }
        .file-drop-zone .icon { font-size: 2rem; color: var(--gray-400); display: block; margin-bottom: 0.5rem; }
        .file-drop-zone .file-name { color: var(--primary); font-weight: 500; }
        .file-drop-zone input[type="file"] { display: none; }
        
        /* Viewer */
        .viewer-container { background: white; border-radius: var(--radius-lg); box-shadow: var(--shadow-sm); padding: 2rem; margin-bottom: 2rem; border: 1px solid var(--gray-200); }
        .viewer-header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 1.5rem; }
        .viewer-header h2 { font-size: 1.25rem; color: var(--gray-900); }
        .viewer-header .reference { font-size: 0.85rem; color: var(--gray-400); font-family: monospace; }
        .viewer-body { min-height: 300px; background: var(--gray-50); border-radius: var(--radius); padding: 1.5rem; }
        .viewer-body .pdf-viewer { width: 100%; height: 600px; border: none; border-radius: var(--radius); }
        .viewer-body .image-viewer { max-width: 100%; max-height: 600px; display: block; margin: 0 auto; border-radius: var(--radius); }
        .viewer-body .text-viewer { background: white; padding: 1.5rem; border-radius: var(--radius); white-space: pre-wrap; font-size: 0.9rem; line-height: 1.8; max-height: 600px; overflow-y: auto; }
        .viewer-body .empty-viewer { text-align: center; padding: 3rem; color: var(--gray-400); }
        .viewer-body .empty-viewer i { font-size: 3rem; display: block; margin-bottom: 1rem; }
        
        /* Toolbar */
        .toolbar { display: flex; flex-wrap: wrap; gap: 0.75rem; margin-bottom: 2rem; align-items: center; }
        .toolbar .filter-group { display: flex; gap: 0.5rem; flex-wrap: wrap; flex: 1; }
        .toolbar .filter-group select { padding: 0.5rem 1rem; border: 1px solid var(--gray-200); border-radius: var(--radius); background: white; font-size: 0.85rem; color: var(--gray-700); transition: var(--transition); cursor: pointer; }
        .toolbar .filter-group select:focus { outline: none; border-color: var(--primary); box-shadow: 0 0 0 3px rgba(214, 28, 78, 0.1); }
        .toolbar .actions { display: flex; gap: 0.5rem; flex-wrap: wrap; }
        
        /* Toast */
        .toast-container { position: fixed; bottom: 2rem; right: 2rem; z-index: 9999; display: flex; flex-direction: column; gap: 0.5rem; max-width: 400px; }
        .toast { padding: 0.75rem 1.25rem; border-radius: var(--radius); background: white; box-shadow: var(--shadow-lg); display: flex; align-items: center; gap: 0.75rem; animation: slideIn 0.3s ease; border-left: 4px solid var(--gray-400); }
        .toast.success { border-left-color: var(--success); }
        .toast.error { border-left-color: var(--danger); }
        .toast.warning { border-left-color: var(--warning); }
        .toast.info { border-left-color: var(--accent); }
        .toast i { font-size: 1.2rem; }
        .toast.success i { color: var(--success); }
        .toast.error i { color: var(--danger); }
        .toast.warning i { color: var(--warning); }
        .toast.info i { color: var(--accent); }
        .toast .message { flex: 1; font-size: 0.9rem; }
        .toast .close { background: transparent; border: none; color: var(--gray-400); cursor: pointer; padding: 0.25rem; font-size: 1rem; }
        @keyframes slideIn { from { opacity: 0; transform: translateX(20px); } to { opacity: 1; transform: translateX(0); } }
        
        /* Responsive */
        @media (max-width: 1024px) { .nav-container { gap: 1rem; } }
        @media (max-width: 768px) {
            .navbar { padding: 0 1rem; }
            .nav-search { max-width: none; }
            .nav-brand .logo-text .sub { display: none; }
            .user-info { display: none; }
            .notif-panel { width: calc(100% - 2rem); right: 1rem; top: 68px; }
            .main-content { padding: 1rem; }
            .fiches-grid { grid-template-columns: 1fr; }
            .form-row { grid-template-columns: 1fr; }
            .toolbar { flex-direction: column; align-items: stretch; }
            .toolbar .filter-group { flex-direction: column; }
            .stats-bar { grid-template-columns: 1fr 1fr; }
            .viewer-body .pdf-viewer { height: 400px; }
        }
        @media (max-width: 480px) {
            .stats-bar { grid-template-columns: 1fr; }
            .nav-right { gap: 0.5rem; }
            .modal { padding: 0; }
            .modal-body { padding: 1rem; }
        }
    </style>
    {% block extra_head %}{% endblock %}
</head>
<body>
    {% if user.is_authenticated %}
    <nav class="navbar">
        <div class="nav-container">
            <div class="nav-brand">
                <div class="logo-icon"><img src="/static/logo.svg.jpg" alt="Quantum Technology"></div>
                <div class="logo-text">
                    <span class="brand">Quantum</span>
                    <span class="sub">Technology</span>
                </div>
            </div>
            
            <div class="nav-search">
                <form onsubmit="event.preventDefault(); searchFiches();">
                    <input type="text" id="searchInput" placeholder="Rechercher une fiche...">
                    <button type="submit"><i class="fas fa-search"></i></button>
                </form>
            </div>
            
            <div class="nav-right">
                <button class="nav-btn" id="notifBtn" onclick="toggleNotifications()">
                    <i class="fas fa-bell"></i>
                    <span class="notif-badge" id="notifCount">0</span>
                </button>
                
                <div class="user-menu" onclick="toggleDropdown()">
                    <div class="user-avatar">{{ user.username|first|upper }}</div>
                    <div class="user-info">
                        <span class="user-name">{{ user.get_full_name|default:user.username }}</span>
                        <span class="user-role">{% if user.is_staff %}Administrateur{% else %}Utilisateur{% endif %}</span>
                    </div>
                    <i class="fas fa-chevron-down" style="color: var(--gray-400); font-size: 0.8rem;"></i>
                    <div class="dropdown" id="userDropdown">
                        <a href="{% url 'dashboard' %}"><i class="fas fa-tachometer-alt"></i> Tableau de bord</a>
                        <hr>
                        <a href="{% url 'logout' %}" class="text-danger"><i class="fas fa-sign-out-alt"></i> Déconnexion</a>
                    </div>
                </div>
            </div>
        </div>
    </nav>

    <section class="hero-banner">
        <div class="hero-content">
            <div>
                <h1>Accélérez la gestion de vos fiches techniques</h1>
                <p>Centralisez, recherchez, versionnez et visualisez vos documents dans un espace sécurisé à l’identité Quantum.</p>
            </div>
            <button class="btn btn-primary" onclick="openImportModal()"><i class="fas fa-upload"></i> Importer une fiche</button>
        </div>
    </section>
    
    <div class="notif-panel" id="notifPanel">
        <div class="notif-header">
            <h3><i class="fas fa-bell"></i> Notifications</h3>
            <button class="btn-text" onclick="markAllRead()"><i class="fas fa-check-double"></i> Tout marquer lu</button>
        </div>
        <div class="notif-list" id="notifList">
            <div class="notif-empty"><i class="fas fa-inbox"></i><p>Aucune notification</p></div>
        </div>
    </div>
    {% endif %}
    
    <main class="main-content">
        {% block content %}{% endblock %}
    </main>
    
    <div class="toast-container" id="toastContainer"></div>
    
    <script>
    function getCSRFToken() { return document.querySelector('[name=csrfmiddlewaretoken]')?.value || ''; }
    
    function showToast(message, type = 'info', duration = 4000) {
        const container = document.getElementById('toastContainer');
        if (!container) return;
        const icons = { success: 'fa-check-circle', error: 'fa-exclamation-circle', warning: 'fa-exclamation-triangle', info: 'fa-info-circle' };
        const toast = document.createElement('div');
        toast.className = `toast ${type}`;
        toast.innerHTML = `<i class="fas ${icons[type] || icons.info}"></i><span class="message">${message}</span><button class="close" onclick="this.parentElement.remove()"><i class="fas fa-times"></i></button>`;
        container.appendChild(toast);
        setTimeout(() => { toast.style.opacity = '0'; toast.style.transform = 'translateX(20px)'; setTimeout(() => toast.remove(), 300); }, duration);
    }
    
    function formatDate(dateStr) {
        const date = new Date(dateStr);
        const now = new Date();
        const diff = now - date;
        if (diff < 60000) return "À l'instant";
        if (diff < 3600000) return `${Math.floor(diff / 60000)} min`;
        if (diff < 86400000) return `${Math.floor(diff / 3600000)} h`;
        if (diff < 604800000) return `${Math.floor(diff / 86400000)} j`;
        return date.toLocaleDateString('fr-FR', { day: '2-digit', month: '2-digit', year: 'numeric', hour: '2-digit', minute: '2-digit' });
    }
    
    function escapeHtml(text) { if (!text) return ''; const div = document.createElement('div'); div.textContent = text; return div.innerHTML; }
    
    function toggleDropdown() {
        document.getElementById('userDropdown').classList.toggle('active');
    }
    document.addEventListener('click', function(e) {
        const menu = document.querySelector('.user-menu');
        if (menu && !menu.contains(e.target)) { document.getElementById('userDropdown')?.classList.remove('active'); }
    });
    
    function toggleNotifications() {
        const panel = document.getElementById('notifPanel');
        panel.classList.toggle('active');
        if (panel.classList.contains('active')) loadNotifications();
    }
    
    async function loadNotifications() {
        try {
            const response = await fetch('/api/notifications/');
            const data = await response.json();
            const list = document.getElementById('notifList');
            const count = document.getElementById('notifCount');
            const unread = data.filter(n => !n.is_read);
            if (count) count.textContent = unread.length;
            if (!list) return;
            if (data.length === 0) {
                list.innerHTML = '<div class="notif-empty"><i class="fas fa-inbox"></i><p>Aucune notification</p></div>';
                return;
            }
            list.innerHTML = data.map(n => `
                <div class="notif-item ${n.is_read ? '' : 'unread'}">
                    <div class="notif-icon ${n.type}">
                        <i class="fas ${n.type === 'IMPORT' ? 'fa-upload' : n.type === 'UPDATE' ? 'fa-edit' : n.type === 'DELETE' ? 'fa-trash' : 'fa-code-branch'}"></i>
                    </div>
                    <div class="notif-content">
                        <div class="message">${escapeHtml(n.message)}</div>
                        <div class="time">${formatDate(n.created_at)}</div>
                    </div>
                </div>
            `).join('');
        } catch (error) { console.error('Erreur notifications:', error); }
    }
    
    async function markAllRead() {
        try {
            await fetch('/api/notifications/read-all/', { method: 'PUT', headers: { 'X-CSRFToken': getCSRFToken() } });
            loadNotifications();
            showToast('Toutes les notifications marquées comme lues', 'success');
        } catch (error) { showToast('Erreur', 'error'); }
    }
    
    function searchFiches() {
        const query = document.getElementById('searchInput')?.value || '';
        if (typeof loadFiches === 'function') loadFiches();
    }
    
    document.addEventListener('DOMContentLoaded', function() {
        const searchInput = document.getElementById('searchInput');
        if (searchInput) {
            let timeout;
            searchInput.addEventListener('input', function() {
                clearTimeout(timeout);
                timeout = setTimeout(() => { if (typeof loadFiches === 'function') loadFiches(); }, 300);
            });
        }
    });
    
    function openModal(id) {
        const modal = document.getElementById(id);
        if (modal) { modal.classList.add('active'); document.body.style.overflow = 'hidden'; }
    }
    function closeModal(id) {
        const modal = document.getElementById(id);
        if (modal) { modal.classList.remove('active'); document.body.style.overflow = ''; }
    }
    document.addEventListener('keydown', function(e) {
        if (e.key === 'Escape') {
            document.querySelectorAll('.modal-overlay.active').forEach(modal => closeModal(modal.id));
            document.getElementById('notifPanel')?.classList.remove('active');
        }
    });
    document.querySelectorAll('.modal-overlay').forEach(modal => {
        modal.addEventListener('click', function(e) { if (e.target === this) closeModal(this.id); });
    });
    
    document.addEventListener('DOMContentLoaded', function() {
        document.querySelectorAll('.file-drop-zone').forEach(zone => {
            const input = zone.querySelector('input[type="file"]');
            zone.addEventListener('click', () => input?.click());
            zone.addEventListener('dragover', (e) => { e.preventDefault(); zone.classList.add('dragover'); });
            zone.addEventListener('dragleave', () => zone.classList.remove('dragover'));
            zone.addEventListener('drop', (e) => {
                e.preventDefault();
                zone.classList.remove('dragover');
                const files = e.dataTransfer.files;
                if (files.length > 0 && input) {
                    input.files = files;
                    const fileName = zone.querySelector('.file-name');
                    if (fileName) fileName.textContent = files[0].name;
                    const p = zone.querySelector('p');
                    if (p) p.textContent = `Fichier sélectionné : ${files[0].name}`;
                }
            });
            if (input) {
                input.addEventListener('change', function() {
                    if (this.files.length > 0) {
                        const fileName = zone.querySelector('.file-name');
                        if (fileName) fileName.textContent = this.files[0].name;
                        const p = zone.querySelector('p');
                        if (p) p.textContent = `Fichier sélectionné : ${this.files[0].name}`;
                    }
                });
            }
        });
    });
    console.log('Quantum script loaded - version: 20260709-1');
    </script>
    {% block extra_scripts %}{% endblock %}
</body>
</html>
    ''',
    
    'login.html': '''
{% extends 'base.html' %}
{% block title %}Connexion - Quantum Technology{% endblock %}
{% block content %}
<div style="max-width: 420px; margin: 3rem auto;">
    <div style="background: white; border-radius: var(--radius-lg); box-shadow: var(--shadow-lg); padding: 2.5rem; border: 1px solid var(--gray-200);">
        <div style="text-align: center; margin-bottom: 2rem;">
            <div style="width: 64px; height: 64px; background: var(--primary); border-radius: 16px; display: flex; align-items: center; justify-content: center; margin: 0 auto 1rem;">
                <span style="color: white; font-weight: 800; font-size: 1.8rem; letter-spacing: -1px;">Q</span>
            </div>
            <h1 style="font-size: 1.5rem; font-weight: 700; color: var(--gray-900);">Connexion</h1>
            <p style="color: var(--gray-500); font-size: 0.9rem;">Accédez à votre espace de travail</p>
        </div>
        
        <form method="post" action="{% url 'login' %}">
            {% csrf_token %}
            {% if error %}
                <div style="background: rgba(214,28,78,0.12); color: #A31B45; border-radius: 12px; padding: 0.9rem; margin-bottom: 1rem;">
                    {{ error }}
                </div>
            {% endif %}
            <div class="form-group">
                <label>Email</label>
                <input type="email" name="email" placeholder="votre@email.com" required>
            </div>
            <div class="form-group">
                <label>Mot de passe</label>
                <input type="password" name="password" placeholder="Entrez votre mot de passe" required>
            </div>
            <button type="submit" class="btn btn-primary" style="width: 100%; justify-content: center; padding: 0.75rem;">
                <i class="fas fa-sign-in-alt"></i> Se connecter
            </button>
        </form>
        
        <div style="text-align: center; margin-top: 1.5rem; padding-top: 1.5rem; border-top: 1px solid var(--gray-200);">
            <p style="color: var(--gray-500); font-size: 0.9rem;">
                Pas encore de compte ? 
                <a href="{% url 'register' %}" style="color: var(--primary); font-weight: 500; text-decoration: none;">Créer un compte</a>
            </p>
        </div>
    </div>
</div>
{% endblock %}
    ''',
    
    'register.html': '''
{% extends 'base.html' %}
{% block title %}Inscription - Quantum Technology{% endblock %}
{% block content %}
<div style="max-width: 420px; margin: 3rem auto;">
    <div style="background: white; border-radius: var(--radius-lg); box-shadow: var(--shadow-lg); padding: 2.5rem; border: 1px solid var(--gray-200);">
        <div style="text-align: center; margin-bottom: 2rem;">
            <div style="width: 96px; height: 96px; background: transparent; border-radius: 20px; display: flex; align-items: center; justify-content: center; margin: 0 auto 1.25rem; overflow: hidden;">
                <img src="/static/logo.svg.jpg" alt="Quantum Technology" style="width: 100%; height: 100%; object-fit: contain; display: block;" />
            </div>
            <h1 style="font-size: 1.5rem; font-weight: 700; color: var(--gray-900);">Créer un compte</h1>
            <p style="color: var(--gray-500); font-size: 0.9rem;">Rejoignez Quantum Technology</p>
        </div>
        
        <form method="post" action="{% url 'register' %}">
            {% csrf_token %}
            {% if error %}
                <div style="background: rgba(214,28,78,0.12); color: #A31B45; border-radius: 12px; padding: 0.9rem; margin-bottom: 1rem;">
                    {{ error }}
                </div>
            {% endif %}
            <div class="form-group">
                <label>Nom d'utilisateur <span class="required">*</span></label>
                <input type="text" name="username" placeholder="Choisissez un nom" required>
            </div>
            <div class="form-group">
                <label>Email <span class="required">*</span></label>
                <input type="text" name="email" placeholder="votre@email.com" required>
            </div>
            <div class="form-group">
                <label>Nom complet <span class="required">*</span></label>
                <input type="text" name="full_name" placeholder="Votre nom complet" required>
            </div>
            <div class="form-group">
                <label>Mot de passe <span class="required">*</span></label>
                <input type="password" name="password1" placeholder="Minimum 6 caractères" required minlength="6">
            </div>
            <div class="form-group">
                <label>Confirmer le mot de passe <span class="required">*</span></label>
                <input type="password" name="password2" placeholder="Confirmez votre mot de passe" required>
            </div>
            <button type="submit" class="btn btn-secondary" style="width: 100%; justify-content: center; padding: 0.75rem;">
                <i class="fas fa-user-plus"></i> S'inscrire
            </button>
        </form>
        
        <div style="text-align: center; margin-top: 1.5rem; padding-top: 1.5rem; border-top: 1px solid var(--gray-200);">
            <p style="color: var(--gray-500); font-size: 0.9rem;">
                Déjà un compte ? 
                <a href="{% url 'login' %}" style="color: var(--primary); font-weight: 500; text-decoration: none;">Se connecter</a>
            </p>
        </div>
    </div>
</div>
{% endblock %}
    ''',
    
    'dashboard.html': '''
{% extends 'base.html' %}
{% block title %}Tableau de bord - Quantum Technology{% endblock %}
{% block content %}
<div class="stats-bar" id="statsBar">
    <div class="stat-card">
        <div class="stat-header">
            <span class="label">Fiches techniques</span>
            <div class="icon primary"><i class="fas fa-file-alt"></i></div>
        </div>
        <div class="stat-value" id="statTotal">0</div>
        <div class="stat-footer">Total des fiches importées</div>
    </div>
    <div class="stat-card">
        <div class="stat-header">
            <span class="label">Catégories</span>
            <div class="icon secondary"><i class="fas fa-tags"></i></div>
        </div>
        <div class="stat-value" id="statCategories">0</div>
        <div class="stat-footer">Catégories disponibles</div>
    </div>
    <div class="stat-card">
        <div class="stat-header">
            <span class="label">Tags</span>
            <div class="icon success"><i class="fas fa-hashtag"></i></div>
        </div>
        <div class="stat-value" id="statTags">0</div>
        <div class="stat-footer">Tags utilisés</div>
    </div>
    <div class="stat-card">
        <div class="stat-header">
            <span class="label">Versions</span>
            <div class="icon accent"><i class="fas fa-code-branch"></i></div>
        </div>
        <div class="stat-value" id="statVersions">0</div>
        <div class="stat-footer">Versions totales</div>
    </div>
</div>

<div class="toolbar">
    <div class="filter-group">
        <select id="categoryFilter" onchange="loadFiches()">
            <option value="">Toutes les catégories</option>
            <option value="ELECTRONIQUE">Électronique</option>
            <option value="MECANIQUE">Mécanique</option>
            <option value="LOGICIEL">Logiciel</option>
            <option value="MATERIAUX">Matériaux</option>
            <option value="ELECTRIQUE">Électrique</option>
            <option value="AUTRE">Autre</option>
        </select>
        <select id="tagFilter" onchange="loadFiches()">
            <option value="">Tous les tags</option>
        </select>
    </div>
    <div class="actions">
        <button class="btn btn-primary" onclick="loadFiches()"><i class="fas fa-sync-alt"></i> Actualiser</button>
        <button class="btn btn-secondary" onclick="openImportModal()"><i class="fas fa-upload"></i> Importer</button>
    </div>
</div>

<div id="fichesGrid" class="fiches-grid">
    <div class="loading-spinner"><div class="spinner"></div></div>
</div>

<!-- Import Modal -->
<div id="importModal" class="modal-overlay">
    <div class="modal">
        <div class="modal-header">
            <h2><i class="fas fa-upload"></i> Importer une fiche technique</h2>
            <button class="close-btn" onclick="closeModal('importModal')">&times;</button>
        </div>
        <form id="importForm" enctype="multipart/form-data">
            {% csrf_token %}
            <div class="modal-body">
                <div class="form-group">
                    <label>Fichier <span class="required">*</span></label>
                    <div class="file-drop-zone">
                        <i class="fas fa-cloud-upload-alt icon"></i>
                        <p>Glissez-déposez votre fichier ou cliquez pour sélectionner</p>
                        <span class="file-name"></span>
                        <input type="file" name="file" accept=".pdf,.doc,.docx,.xls,.xlsx,.txt,.jpg,.png" required>
                    </div>
                </div>
                <div class="form-row">
                    <div class="form-group">
                        <label>Référence</label>
                        <input type="text" name="reference" placeholder="FT-2026-001">
                    </div>
                    <div class="form-group">
                        <label>Nom <span class="required">*</span></label>
                        <input type="text" name="name" placeholder="Nom de la fiche" required>
                    </div>
                </div>
                <div class="form-group">
                    <label>Description</label>
                    <textarea name="description" placeholder="Description du produit ou composant"></textarea>
                </div>
                <div class="form-row">
                    <div class="form-group">
                        <label>Catégorie</label>
                        <select name="category">
                            <option value="">Non catégorisé</option>
                            <option value="ELECTRONIQUE">Électronique</option>
                            <option value="MECANIQUE">Mécanique</option>
                            <option value="LOGICIEL">Logiciel</option>
                            <option value="MATERIAUX">Matériaux</option>
                            <option value="ELECTRIQUE">Électrique</option>
                            <option value="AUTRE">Autre</option>
                        </select>
                    </div>
                    <div class="form-group">
                        <label>Fabricant</label>
                        <input type="text" name="manufacturer" placeholder="Nom du fabricant">
                    </div>
                </div>
                <div class="form-group">
                    <label>Tags (séparés par des virgules)</label>
                    <input type="text" name="tags" placeholder="ex: composant, certification">
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-outline" onclick="closeModal('importModal')">Annuler</button>
                <button type="submit" class="btn btn-secondary"><i class="fas fa-upload"></i> Importer</button>
            </div>
        </form>
    </div>
</div>

<!-- Version Modal -->
<div id="versionModal" class="modal-overlay">
    <div class="modal">
        <div class="modal-header">
            <h2><i class="fas fa-code-branch"></i> Ajouter une version</h2>
            <button class="close-btn" onclick="closeModal('versionModal')">&times;</button>
        </div>
        <form id="versionForm" enctype="multipart/form-data">
            {% csrf_token %}
            <input type="hidden" id="versionFicheId" name="fiche_id">
            <div class="modal-body">
                <div class="form-group">
                    <label>Nouveau fichier <span class="required">*</span></label>
                    <div class="file-drop-zone">
                        <i class="fas fa-cloud-upload-alt icon"></i>
                        <p>Cliquez pour sélectionner le fichier</p>
                        <span class="file-name"></span>
                        <input type="file" name="file" accept=".pdf,.doc,.docx,.xls,.xlsx,.txt" required>
                    </div>
                </div>
                <div class="form-group">
                    <label>Commentaire</label>
                    <textarea name="comment" placeholder="Description des changements"></textarea>
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-outline" onclick="closeModal('versionModal')">Annuler</button>
                <button type="submit" class="btn btn-primary"><i class="fas fa-plus"></i> Ajouter</button>
            </div>
        </form>
    </div>
</div>

<!-- History Modal -->
<div id="historyModal" class="modal-overlay">
    <div class="modal">
        <div class="modal-header">
            <h2><i class="fas fa-history"></i> Historique - <span id="historyFicheName"></span></h2>
            <button class="close-btn" onclick="closeModal('historyModal')">&times;</button>
        </div>
        <div class="modal-body" id="historyList">
            <div class="loading-spinner"><div class="spinner"></div></div>
        </div>
        <div class="modal-footer">
            <button class="btn btn-outline" onclick="closeModal('historyModal')">Fermer</button>
        </div>
    </div>
</div>

<!-- Edit Modal -->
<div id="editModal" class="modal-overlay">
    <div class="modal">
        <div class="modal-header">
            <h2><i class="fas fa-edit"></i> Modifier la fiche</h2>
            <button class="close-btn" onclick="closeModal('editModal')">&times;</button>
        </div>
        <form id="editForm">
            {% csrf_token %}
            <input type="hidden" id="editId">
            <div class="modal-body">
                <div class="form-group">
                    <label>Nom <span class="required">*</span></label>
                    <input type="text" id="editName" required>
                </div>
                <div class="form-group">
                    <label>Description</label>
                    <textarea id="editDescription"></textarea>
                </div>
                <div class="form-row">
                    <div class="form-group">
                        <label>Catégorie</label>
                        <select id="editCategory">
                            <option value="">Non catégorisé</option>
                            <option value="ELECTRONIQUE">Électronique</option>
                            <option value="MECANIQUE">Mécanique</option>
                            <option value="LOGICIEL">Logiciel</option>
                            <option value="MATERIAUX">Matériaux</option>
                            <option value="ELECTRIQUE">Électrique</option>
                            <option value="AUTRE">Autre</option>
                        </select>
                    </div>
                    <div class="form-group">
                        <label>Fabricant</label>
                        <input type="text" id="editManufacturer" placeholder="Nom du fabricant">
                    </div>
                </div>
                <div class="form-group">
                    <label>Tags (séparés par des virgules)</label>
                    <input type="text" id="editTags" placeholder="ex: composant, certification">
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-outline" onclick="closeModal('editModal')">Annuler</button>
                <button type="submit" class="btn btn-primary"><i class="fas fa-save"></i> Enregistrer</button>
            </div>
        </form>
    </div>
</div>

<!-- View Modal -->
<div id="viewModal" class="modal-overlay">
    <div class="modal" style="max-width: 800px;">
        <div class="modal-header">
            <h2><i class="fas fa-eye"></i> Visualisation - <span id="viewFicheName"></span></h2>
            <button class="close-btn" onclick="closeModal('viewModal')">&times;</button>
        </div>
        <div class="modal-body" id="viewContent">
            <div class="loading-spinner"><div class="spinner"></div></div>
        </div>
        <div class="modal-footer">
            <button class="btn btn-secondary" id="viewDownloadBtn"><i class="fas fa-download"></i> Télécharger</button>
            <button class="btn btn-outline" onclick="closeModal('viewModal')">Fermer</button>
        </div>
    </div>
</div>

<script>
async function loadFiches() {
    const search = document.getElementById('searchInput')?.value || '';
    const category = document.getElementById('categoryFilter')?.value || '';
    const tag = document.getElementById('tagFilter')?.value || '';
    const grid = document.getElementById('fichesGrid');
    if (grid) grid.innerHTML = '<div class="loading-spinner"><div class="spinner"></div></div>';
    try {
        let url = '/api/fiches/?';
        if (search) url += `search=${encodeURIComponent(search)}&`;
        if (category) url += `category=${encodeURIComponent(category)}&`;
        if (tag) url += `tag=${encodeURIComponent(tag)}&`;
        const response = await fetch(url);
        const data = await response.json();
        if (!response.ok) throw new Error(data.error || 'Erreur de chargement');
        const fiches = Array.isArray(data.results) ? data.results : Array.isArray(data) ? data : [];
        renderFiches(fiches);
    } catch (error) {
        console.error('Erreur:', error);
        if (grid) {
            grid.innerHTML = `<div class="empty-state"><i class="fas fa-exclamation-circle"></i><h3>Erreur de chargement</h3><p>Impossible de charger les fiches.</p><button class="btn btn-primary" onclick="loadFiches()"><i class="fas fa-sync"></i> Réessayer</button></div>`;
        }
    }
}

function renderFiches(fiches) {
    const grid = document.getElementById('fichesGrid');
    if (!grid) return;
    if (!fiches || fiches.length === 0) {
        grid.innerHTML = `<div class="empty-state"><i class="fas fa-file-alt"></i><h3>Aucune fiche trouvée</h3><p>Commencez par importer votre première fiche technique.</p><button class="btn btn-secondary" onclick="openImportModal()"><i class="fas fa-upload"></i> Importer</button></div>`;
        return;
    }
    grid.innerHTML = fiches.map(f => `
        <div class="fiche-card">
            <div class="card-header">
                <div class="title">
                    <h3>${escapeHtml(f.name)}</h3>
                    <div class="reference">${escapeHtml(f.reference)}</div>
                </div>
                <span class="version-badge">v${escapeHtml(f.version)}</span>
            </div>
            <div class="card-body">
                <div class="meta"><i class="fas fa-tag"></i> ${f.category_display || 'Non catégorisé'} <span style="margin: 0 0.5rem;">|</span> <i class="fas fa-building"></i> ${escapeHtml(f.manufacturer || 'Fabricant inconnu')}</div>
                ${f.description ? `<div class="description">${escapeHtml(f.description)}</div>` : ''}
                ${f.tags && f.tags.length > 0 ? `<div class="tags">${f.tags.map(t => `<span class="tag" style="background: ${t.color || '#293462'}">${escapeHtml(t.name)}</span>`).join('')}</div>` : ''}
            </div>
            <div class="card-actions">
                <button class="btn btn-accent btn-sm" onclick="viewFiche('${f.id}')"><i class="fas fa-eye"></i> Voir</button>
                <button class="btn btn-primary btn-sm" onclick="downloadFiche('${f.id}')"><i class="fas fa-download"></i></button>
                <button class="btn btn-outline btn-sm" onclick="openVersionModal('${f.id}')"><i class="fas fa-code-branch"></i></button>
                <button class="btn btn-outline btn-sm" onclick="editFiche('${f.id}')"><i class="fas fa-edit"></i></button>
                ${f.versions_count > 0 ? `<button class="btn btn-outline btn-sm" onclick="showHistory('${f.id}')"><i class="fas fa-history"></i> ${f.versions_count}</button>` : ''}
                <button class="btn btn-danger btn-sm" onclick="deleteFiche('${f.id}')"><i class="fas fa-trash"></i></button>
            </div>
            <div class="card-footer">
                <span class="author"><i class="fas fa-user"></i> ${escapeHtml(f.author_display || f.author || 'Inconnu')}</span>
                <span class="date"><i class="fas fa-calendar-alt"></i> ${formatDate(f.created_at)}</span>
            </div>
        </div>
    `).join('');
}

function openImportModal() {
    document.getElementById('importForm').reset();
    document.querySelector('#importModal .file-name').textContent = '';
    openModal('importModal');
}

document.getElementById('importForm')?.addEventListener('submit', async function(e) {
    e.preventDefault();
    const formData = new FormData(this);
    const submitBtn = this.querySelector('button[type="submit"]');
    const originalText = submitBtn?.innerHTML;
    if (submitBtn) { submitBtn.innerHTML = '<i class="fas fa-spinner fa-spin"></i> Importation...'; submitBtn.disabled = true; }
    try {
        const response = await fetch('/api/fiches/create/', { method: 'POST', headers: { 'X-CSRFToken': getCSRFToken() }, body: formData });
        const data = await response.json();
        if (!response.ok) throw new Error(data.error || 'Erreur');
        showToast('Fiche importée avec succès', 'success');
        closeModal('importModal');
        loadFiches();
        loadStats();
        loadTags();
    } catch (error) { showToast(error.message || 'Erreur', 'error'); }
    finally { if (submitBtn) { submitBtn.innerHTML = originalText; submitBtn.disabled = false; } }
});

async function downloadFiche(id) { window.open(`/api/fiches/${id}/download/`, '_blank'); }

async function deleteFiche(id) {
    if (!confirm('Supprimer définitivement cette fiche ?')) return;
    try {
        const response = await fetch(`/api/fiches/${id}/delete/`, { method: 'DELETE', headers: { 'X-CSRFToken': getCSRFToken() } });
        if (!response.ok) throw new Error('Erreur');
        showToast('Fiche supprimée', 'success');
        loadFiches();
        loadStats();
    } catch (error) { showToast('Erreur lors de la suppression', 'error'); }
}

async function viewFiche(id) {
    const modal = document.getElementById('viewModal');
    const content = document.getElementById('viewContent');
    const name = document.getElementById('viewFicheName');
    const downloadBtn = document.getElementById('viewDownloadBtn');
    if (!modal) return;
    openModal('viewModal');
    content.innerHTML = '<div class="loading-spinner"><div class="spinner"></div></div>';
    try {
        const response = await fetch(`/api/fiches/${id}/view/`);
        const data = await response.json();
        if (name) name.textContent = data.name || `Fiche ${data.reference}`;
        if (downloadBtn) downloadBtn.onclick = () => downloadFiche(id);
        const preview = data.preview || {};
        let html = '';
        if (data.file_type && data.file_type.includes('pdf')) {
            html = `<div class="viewer-body"><iframe class="pdf-viewer" src="${data.file_url}#toolbar=0"></iframe><div style="text-align: center; margin-top: 0.5rem; color: var(--gray-400); font-size: 0.8rem;"><i class="fas fa-file-pdf"></i> PDF - ${preview.pages || '?'} pages</div></div>`;
        } else if (data.file_type && data.file_type.includes('image')) {
            html = `<div class="viewer-body"><img class="image-viewer" src="${data.file_url}" alt="${data.name}"></div>`;
        } else if (preview.text) {
            html = `<div class="viewer-body"><div class="text-viewer">${escapeHtml(preview.text)}</div></div>`;
        } else {
            html = `<div class="viewer-body"><div class="empty-viewer"><i class="fas fa-file"></i><p>Aucun aperçu disponible</p><button class="btn btn-secondary" onclick="downloadFiche('${id}')" style="margin-top: 0.5rem;"><i class="fas fa-download"></i> Télécharger</button></div></div>`;
        }
        content.innerHTML = html;
    } catch (error) {
        content.innerHTML = `<div class="viewer-body"><div class="empty-viewer"><i class="fas fa-exclamation-circle"></i><p>Erreur de chargement</p><button class="btn btn-secondary" onclick="downloadFiche('${id}')" style="margin-top: 0.5rem;"><i class="fas fa-download"></i> Télécharger</button></div></div>`;
    }
}

function openVersionModal(ficheId) {
    document.getElementById('versionFicheId').value = ficheId;
    document.getElementById('versionForm').reset();
    document.querySelector('#versionModal .file-name').textContent = '';
    openModal('versionModal');
}

document.getElementById('versionForm')?.addEventListener('submit', async function(e) {
    e.preventDefault();
    const ficheId = document.getElementById('versionFicheId').value;
    const formData = new FormData(this);
    const submitBtn = this.querySelector('button[type="submit"]');
    const originalText = submitBtn?.innerHTML;
    if (submitBtn) { submitBtn.innerHTML = '<i class="fas fa-spinner fa-spin"></i> Ajout...'; submitBtn.disabled = true; }
    try {
        const response = await fetch(`/api/fiches/${ficheId}/version/`, { method: 'POST', headers: { 'X-CSRFToken': getCSRFToken() }, body: formData });
        const data = await response.json();
        if (!response.ok) throw new Error(data.error || 'Erreur');
        showToast('Version ajoutée avec succès', 'success');
        closeModal('versionModal');
        loadFiches();
    } catch (error) { showToast(error.message || 'Erreur', 'error'); }
    finally { if (submitBtn) { submitBtn.innerHTML = originalText; submitBtn.disabled = false; } }
});

async function editFiche(id) {
    try {
        const response = await fetch(`/api/fiches/${id}/`);
        if (!response.ok) {
            const err = await response.json().catch(() => ({}));
            throw new Error(err.error || 'Erreur de chargement');
        }
        const data = await response.json();
        document.getElementById('editId').value = data.id;
        document.getElementById('editName').value = data.name;
        document.getElementById('editDescription').value = data.description || '';
        document.getElementById('editCategory').value = data.category || '';
        document.getElementById('editManufacturer').value = data.manufacturer || '';
        document.getElementById('editTags').value = (data.tags || []).map(t => t.name || t).join(', ');
        openModal('editModal');
    } catch (error) { showToast('Erreur de chargement', 'error'); }
}

document.getElementById('editForm')?.addEventListener('submit', async function(e) {
    e.preventDefault();
    const id = document.getElementById('editId').value;
    const data = {
        name: document.getElementById('editName').value,
        description: document.getElementById('editDescription').value,
        category: document.getElementById('editCategory').value,
        manufacturer: document.getElementById('editManufacturer').value,
        tags: document.getElementById('editTags').value.split(',').map(t => t.trim()).filter(t => t)
    };
    const submitBtn = this.querySelector('button[type="submit"]');
    const originalText = submitBtn?.innerHTML;
    if (submitBtn) { submitBtn.innerHTML = '<i class="fas fa-spinner fa-spin"></i> Enregistrement...'; submitBtn.disabled = true; }
    try {
        const response = await fetch(`/api/fiches/${id}/update/`, { method: 'PUT', headers: { 'Content-Type': 'application/json', 'X-CSRFToken': getCSRFToken() }, body: JSON.stringify(data) });
        if (!response.ok) throw new Error('Erreur');
        showToast('Fiche mise à jour', 'success');
        closeModal('editModal');
        loadFiches();
    } catch (error) { showToast('Erreur', 'error'); }
    finally { if (submitBtn) { submitBtn.innerHTML = originalText; submitBtn.disabled = false; } }
});

async function showHistory(id) {
    try {
        const response = await fetch(`/api/fiches/${id}/`);
        const data = await response.json();
        const list = document.getElementById('historyList');
        const name = document.getElementById('historyFicheName');
        if (name) name.textContent = data.name;
        if (!data.versions || data.versions.length === 0) {
            list.innerHTML = '<div class="empty-state"><i class="fas fa-history"></i><p>Aucune version disponible</p></div>';
        } else {
            list.innerHTML = data.versions.map(v => `
                <div style="display: flex; justify-content: space-between; align-items: center; padding: 0.75rem; border-bottom: 1px solid var(--gray-200);">
                    <div>
                        <span style="font-weight: 600; color: var(--gray-800);">v${v.version}</span>
                        ${v.comment ? `<p style="font-size: 0.85rem; color: var(--gray-500); margin-top: 0.25rem;">${escapeHtml(v.comment)}</p>` : ''}
                        <span style="font-size: 0.75rem; color: var(--gray-400);">${formatDate(v.created_at)} - ${v.author || 'Inconnu'}</span>
                    </div>
                    <a href="/api/fiches/${id}/versions/${v.id}/download/" target="_blank" class="btn btn-accent btn-sm"><i class="fas fa-download"></i></a>
                </div>
            `).join('');
        }
        openModal('historyModal');
    } catch (error) { showToast('Erreur de chargement', 'error'); }
}

async function loadStats() {
    try {
        const response = await fetch('/api/stats/');
        const data = await response.json();
        document.getElementById('statTotal').textContent = data.total_fiches || 0;
        document.getElementById('statCategories').textContent = data.categories_count || 0;
        document.getElementById('statTags').textContent = data.tags_count || 0;
        document.getElementById('statVersions').textContent = data.versions_count || 0;
    } catch (error) { console.error('Erreur stats:', error); }
}

async function loadTags() {
    try {
        const response = await fetch('/api/tags/');
        const data = await response.json();
        const select = document.getElementById('tagFilter');
        if (!select) return;
        const currentValue = select.value;
        select.innerHTML = '<option value="">Tous les tags</option>';
        data.forEach(tag => {
            const option = document.createElement('option');
            option.value = tag.name;
            option.textContent = tag.name;
            select.appendChild(option);
        });
        select.value = currentValue;
    } catch (error) { console.error('Erreur tags:', error); }
}

function initDashboard() {
    if (window.dashboardInitialized) return;
    window.dashboardInitialized = true;
    console.debug('Initialisation du tableau de bord');
    loadFiches();
    loadStats();
    loadTags();
    loadNotifications();
    setInterval(loadNotifications, 30000);
}

if (document.readyState === 'loading') {
    document.addEventListener('DOMContentLoaded', initDashboard);
} else {
    initDashboard();
}

// Expose dashboard actions globally for inline onclick handlers
window.loadFiches = loadFiches;
window.openImportModal = openImportModal;
window.viewFiche = viewFiche;
window.downloadFiche = downloadFiche;
window.openVersionModal = openVersionModal;
window.editFiche = editFiche;
window.deleteFiche = deleteFiche;
window.showHistory = showHistory;
window.loadStats = loadStats;
window.loadTags = loadTags;
window.searchFiches = searchFiches;
window.initDashboard = initDashboard;
</script>
{% endblock %}
    ''',

    'fiche_detail.html': '''
{% extends 'base.html' %}
{% block title %}{{ fiche.name }} - Quantum Technology{% endblock %}
{% block content %}
<div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 2rem; flex-wrap: wrap; gap: 1rem;">
    <div>
        <h1 style="font-size: 1.5rem; font-weight: 700; color: var(--gray-900);">{{ fiche.name }}</h1>
        <p style="color: var(--gray-500); font-size: 0.9rem;">
            <span style="font-family: monospace;">{{ fiche.reference }}</span>
            <span style="margin: 0 0.5rem;">|</span>
            v{{ fiche.version }}
            <span style="margin: 0 0.5rem;">|</span>
            {{ fiche.category|default:'Non catégorisé' }}
        </p>
    </div>
    <div style="display: flex; gap: 0.75rem;">
        <a href="{% url 'dashboard' %}" class="btn btn-outline"><i class="fas fa-arrow-left"></i> Retour</a>
        <button class="btn btn-primary" onclick="downloadFiche('{{ fiche.id }}')"><i class="fas fa-download"></i> Télécharger</button>
    </div>
</div>

<div class="viewer-container">
    <div class="viewer-header">
        <div>
            <h2><i class="fas fa-eye"></i> Aperçu du document</h2>
            <span class="reference">{{ fiche.file_name }}</span>
        </div>
        <span style="font-size: 0.8rem; color: var(--gray-400);">{{ fiche.file_size|default:'' }}</span>
    </div>
    <div class="viewer-body" id="viewerContent">
        <div class="loading-spinner"><div class="spinner"></div></div>
    </div>
</div>

<div style="display: grid; grid-template-columns: 1fr 1fr; gap: 2rem; margin-bottom: 2rem;">
    <div>
        <h3 style="font-size: 1rem; font-weight: 600; margin-bottom: 0.75rem;">Description</h3>
        <p style="color: var(--gray-600);">{{ fiche.description|default:'Aucune description' }}</p>
    </div>
    <div>
        <h3 style="font-size: 1rem; font-weight: 600; margin-bottom: 0.75rem;">Informations</h3>
        <ul style="list-style: none; color: var(--gray-600);">
            <li style="padding: 0.25rem 0;"><strong>Fabricant :</strong> {{ fiche.manufacturer|default:'Inconnu' }}</li>
            <li style="padding: 0.25rem 0;"><strong>Auteur :</strong> {{ fiche.profiles.username|default:'Inconnu' }}</li>
            <li style="padding: 0.25rem 0;"><strong>Date :</strong> {{ fiche.created_at|date:'d/m/Y H:i' }}</li>
            <li style="padding: 0.25rem 0;">
                <strong>Tags :</strong>
                {% for tag in fiche.tags %}
                    <span class="tag" style="background: {{ tag.color|default:'#293462' }}; padding: 0.15rem 0.6rem; border-radius: 12px; font-size: 0.7rem; color: white;">#{{ tag.name }}</span>
                {% empty %}
                    Aucun tag
                {% endfor %}
            </li>
        </ul>
    </div>
</div>

<div style="background: white; border-radius: var(--radius-lg); padding: 1.5rem; border: 1px solid var(--gray-200);">
    <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 1rem;">
        <h3 style="font-size: 1rem; font-weight: 600;"><i class="fas fa-code-branch"></i> Historique des versions</h3>
        <button class="btn btn-primary btn-sm" onclick="openVersionModal('{{ fiche.id }}')"><i class="fas fa-plus"></i> Ajouter</button>
    </div>
    {% if fiche.versions %}
        <div style="display: flex; flex-direction: column; gap: 0.5rem;">
            {% for version in fiche.versions %}
                <div style="display: flex; justify-content: space-between; align-items: center; padding: 0.75rem; background: var(--gray-50); border-radius: var(--radius);">
                    <div>
                        <span style="font-weight: 600;">v{{ version.version }}</span>
                        <span style="color: var(--gray-500); font-size: 0.85rem;">{{ version.created_at|date:'d/m/Y H:i' }}</span>
                        {% if version.comment %}
                            <p style="font-size: 0.85rem; color: var(--gray-600);">{{ version.comment }}</p>
                        {% endif %}
                    </div>
                    <a href="/api/fiches/{{ fiche.id }}/versions/{{ version.id }}/download/" class="btn btn-accent btn-sm"><i class="fas fa-download"></i></a>
                </div>
            {% endfor %}
        </div>
    {% else %}
        <p style="color: var(--gray-400); text-align: center; padding: 1rem;">Aucune version disponible</p>
    {% endif %}
</div>

<script>
async function loadViewer() {
    const content = document.getElementById('viewerContent');
    try {
        const response = await fetch('/api/fiches/{{ fiche.id }}/view/');
        const data = await response.json();
        const preview = data.preview || {};
        let html = '';
        if (data.file_type && data.file_type.includes('pdf')) {
            html = `<iframe class="pdf-viewer" src="${data.file_url}#toolbar=0"></iframe>`;
        } else if (data.file_type && data.file_type.includes('image')) {
            html = `<img class="image-viewer" src="${data.file_url}" alt="${data.name}">`;
        } else if (preview.text) {
            html = `<div class="text-viewer">${escapeHtml(preview.text)}</div>`;
        } else {
            html = `<div class="empty-viewer"><i class="fas fa-file"></i><p>Aucun aperçu disponible</p><button class="btn btn-secondary" onclick="downloadFiche('{{ fiche.id }}')" style="margin-top: 0.5rem;"><i class="fas fa-download"></i> Télécharger</button></div>`;
        }
        content.innerHTML = html;
    } catch (error) {
        content.innerHTML = `<div class="empty-viewer"><i class="fas fa-exclamation-circle"></i><p>Erreur de chargement</p></div>`;
    }
}

function downloadFiche(id) { window.open(`/api/fiches/${id}/download/`, '_blank'); }

document.addEventListener('DOMContentLoaded', loadViewer);
</script>
{% endblock %}
    '''
}

# =================================================================================
# APPLICATION WSGI
# =================================================================================

def create_app():
    from django.core.wsgi import get_wsgi_application
    
    templates_dir = BASE_DIR / 'templates'
    os.makedirs(templates_dir, exist_ok=True)
    
    for name, content in TEMPLATE_FILES.items():
        with open(templates_dir / name, 'w', encoding='utf-8') as f:
            f.write(content)
    
    return get_wsgi_application()

application = create_app()

# =================================================================================
# URLS
# =================================================================================

urlpatterns = [
    path('', login_view, name='home'),
    path('login/', login_view, name='login'),
    path('register/', register_view, name='register'),
    path('dashboard/', dashboard_view, name='dashboard'),
    path('fiche/<str:fiche_id>/', fiche_detail_view, name='fiche_detail'),
    path('logout/', logout_view, name='logout'),
    
    path('api/auth/login/', api_login, name='api_login'),
    path('api/auth/register/', api_register, name='api_register'),
    
    path('api/fiches/', api_fiches_list, name='api_fiches'),
    path('api/fiches/create/', api_fiche_create, name='api_fiche_create'),
    path('api/fiches/<str:fiche_id>/', api_fiche_detail, name='api_fiche_detail'),
    path('api/fiches/<str:fiche_id>/update/', api_fiche_update, name='api_fiche_update'),
    path('api/fiches/<str:fiche_id>/delete/', api_fiche_delete, name='api_fiche_delete'),
    path('api/fiches/<str:fiche_id>/version/', api_fiche_version, name='api_fiche_version'),
    path('api/fiches/<str:fiche_id>/download/', api_fiche_download, name='api_fiche_download'),
    path('api/fiches/<str:fiche_id>/view/', api_fiche_view, name='api_fiche_view'),
    path('api/fiches/<str:fiche_id>/serve/', api_fiche_serve, name='api_fiche_serve'),
    path('api/fiches/<str:fiche_id>/versions/<str:version_id>/download/', api_version_download, name='api_version_download'),
    
    path('api/tags/', api_tags, name='api_tags'),
    path('api/notifications/', api_notifications, name='api_notifications'),
    path('api/notifications/<str:notification_id>/read/', api_notifications_mark_read, name='api_notifications_mark_read'),
    path('api/notifications/read-all/', api_notifications_mark_all_read, name='api_notifications_mark_all_read'),
    path('api/stats/', api_stats, name='api_stats'),
    path('_debug/test_supabase_insert/', debug_supabase_insert, name='debug_supabase_insert'),
]

if settings.DEBUG:
    urlpatterns += staticfiles_urlpatterns()
    urlpatterns += static(settings.MEDIA_URL, document_root=settings.MEDIA_ROOT)

# =================================================================================
# POINT D'ENTRÉE
# =================================================================================

if __name__ == '__main__':
    print("\n" + "=" * 80)
    print("⚡ QUANTUM TECHNOLOGY - GESTIONNAIRE DE FICHES")
    print("=" * 80)
    print("🎨 Design : #D61C4E (Rouge) + #293462 (Bleu nuit)")
    print("📦 Base de données : Supabase (PostgreSQL)")
    print("📋 Fonctionnalités :")
    print("   ✅ Authentification Supabase Auth")
    print("   ✅ CRUD complet des fiches")
    print("   ✅ Visualisation en ligne (PDF, images, documents)")
    print("   ✅ Versioning des fichiers")
    print("   ✅ Système de tags")
    print("   ✅ Notifications en temps réel")
    print("   ✅ Statistiques")
    print("   ✅ Stockage Supabase Storage")
    print("=" * 80)
    print("🌐 http://localhost:8000")
    print("=" * 80 + "\n")
    
    execute_from_command_line(['manage.py', 'runserver', '0.0.0.0:8000'])